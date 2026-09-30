import sqlite3
import unittest
from pathlib import Path
from io import BytesIO

from recon_memory import (
    assign_row_ids, create_memory, make_row_id, normalize_invoice_number,
    open_uploaded_memory, export_memory, export_exceptions, memory_filename, result_to_run_lines,
    import_decisions, apply_decisions, save_run, validate_memory,
)


class MemoryTests(unittest.TestCase):
    def test_download_name_is_safe_and_includes_client_fy_and_date(self):
        from datetime import date
        self.assertEqual(memory_filename("Acme & Co", "2025 - 2026", date(2025, 5, 1)),
                         "Acme_Co_FY2025-26_memory_2025-05-01.db")

    def test_result_snapshot_creates_one_stable_line_per_side(self):
        import pandas as pd
        result = pd.DataFrame([{
            "GSTIN_BOOKS": "G1", "Invoice Number_BOOKS": "INV/0458",
            "Invoice Date_BOOKS": "2025-04-30", "Taxable Value_BOOKS": 100,
            "GSTIN_GST": "G1", "Invoice Number_GST": "INV-0458",
            "Invoice Date_GST": "2025-04-30", "Taxable Value_GST": 100,
            "Recon_Status": "Matched", "Match_Logic": "Exact",
        }])
        lines = result_to_run_lines(result, "2025-26")
        self.assertEqual([line["side"] for line in lines], ["B", "G"])
        self.assertNotEqual(lines[0]["row_id"], lines[1]["row_id"])

    def test_exception_workbook_has_row_id_and_decision_dropdown(self):
        from openpyxl import load_workbook
        path = Path(__file__).parents[1] / "_test-exceptions.db"
        try:
            create_memory(path, "Client A", "2025-26")
            save_run(path, "run-x", "April", "GSTR2B", [{
                "row_id": "abc123", "side": "B", "gstin": "G1", "inv_no": "1",
                "taxable": 100, "engine_status": "Invoices Not in GSTR-2B",
            }], "2025-05-01")
            wb = load_workbook(BytesIO(export_exceptions(path)))
            ws = wb["Exceptions"]
            self.assertEqual(ws.cell(1, 1).value, "Row ID")
            self.assertEqual(ws.cell(2, 1).value, "abc123")
            self.assertEqual(len(ws.data_validations.dataValidation), 1)
        finally:
            path.unlink(missing_ok=True)

    def test_import_link_applies_both_sides_and_flags_changed_amount(self):
        from openpyxl import Workbook
        path = Path(__file__).parents[1] / "_test-decisions.db"
        try:
            create_memory(path, "Client A", "2025-26")
            book_id = make_row_id("G1", "BOOKS-1", "B", "", "2025-26")
            gst_id = make_row_id("G1", "PORTAL-1", "G", "", "2025-26")
            lines = [
                {"row_id": book_id, "side": "B", "gstin": "G1", "inv_no": "BOOKS-1", "taxable": 100, "engine_status": "Invoices Not in GSTR-2B"},
                {"row_id": gst_id, "side": "G", "gstin": "G1", "inv_no": "PORTAL-1", "taxable": 100, "engine_status": "Invoices Not in Purchase Books"},
            ]
            save_run(path, "run-1", "April", "GSTR2B", lines, "2025-05-01")
            wb = Workbook(); ws = wb.active
            ws.append(["Row ID", "Decision", "Linked To", "Reason"])
            ws.append([book_id, "Link", "PORTAL-1", "Same invoice, typo in books"])
            stream = BytesIO(); wb.save(stream)
            result = import_decisions(path, stream.getvalue())
            self.assertEqual((result["saved"], len(result["unrecognized"])), (1, 0))
            self.assertEqual(apply_decisions(path, "run-1"), {
                book_id: "Matched (manual)", gst_id: "Matched (manual)"})
            lines[0]["taxable"] = 110
            save_run(path, "run-2", "April", "GSTR2B", lines, "2025-05-02")
            self.assertEqual(apply_decisions(path, "run-2")[book_id], "Needs review")
            self.assertEqual(apply_decisions(path, "run-2")[gst_id], "Needs review")
        finally:
            path.unlink(missing_ok=True)

    def test_row_id_is_stable_and_normalizes_invoice_punctuation(self):
        a = make_row_id("27AAAAA0000A1Z5", "INV/0458", "B", "INV", "2025-26")
        b = make_row_id("27aaaaa0000a1z5", "inv-0458", "B", "INV", "2025-26")
        self.assertEqual(a, b)
        self.assertEqual(normalize_invoice_number("inv/0458"), normalize_invoice_number("INV-0458"))

    def test_assign_row_ids_does_not_depend_on_dataframe_index(self):
        import pandas as pd
        df = pd.DataFrame({"GSTIN": ["G1"], "Invoice Number": ["INV/0458"]}, index=[77])
        row_id = assign_row_ids(df, side="B").iloc[0]["Row ID"]
        other = pd.DataFrame({"GSTIN": ["G1"], "Invoice Number": ["INV-0458"]}, index=[1])
        self.assertEqual(row_id, assign_row_ids(other, side="B").iloc[0]["Row ID"])

    def test_memory_roundtrip_and_replace(self):
        tmp = Path(__file__).parents[1]
        path = tmp / "_test-memory-roundtrip.db"
        uploaded = None
        try:
            create_memory(path, "Client A", "2025 - 2026")
            validate_memory(path, "client a", "2025 - 2026")
            rows = [{"row_id": "r1", "side": "B", "gstin": "G1", "inv_no": "1",
                     "taxable": 100, "engine_status": "Not in GSTR-2B"}]
            save_run(path, "run-1", "April", "GSTR2B", rows, "2025-05-01")
            first = export_memory(path)
            uploaded = open_uploaded_memory(first, "Client A", "2025 - 2026", tmp)
            self.assertEqual(validate_memory(uploaded, "Client A", "2025 - 2026")["client"], "Client A")
            save_run(path, "run-2", "April", "GSTR2B", rows, "2025-05-02")
            db = sqlite3.connect(path)
            try:
                self.assertEqual(db.execute("SELECT COUNT(*) FROM runs").fetchone()[0], 1)
                self.assertEqual(db.execute("SELECT COUNT(*) FROM run_lines").fetchone()[0], 1)
            finally:
                db.close()
        finally:
            path.unlink(missing_ok=True)
            if uploaded:
                Path(uploaded).unlink(missing_ok=True)

    def test_mismatched_identity_rejected(self):
        p = Path(__file__).parents[1] / "_test-memory-identity.db"
        try:
            create_memory(p, "A", "2025-26")
            with self.assertRaises(ValueError):
                validate_memory(p, "B", "2025-26")
        finally:
            p.unlink(missing_ok=True)


if __name__ == "__main__":
    unittest.main()

