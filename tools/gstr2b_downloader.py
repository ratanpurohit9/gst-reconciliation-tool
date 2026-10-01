"""
E-Way Bill & GSTR-2B Downloader (with Smart QRMP/Quarterly Handling, JSON Download & Excel Merge)
-------------------------------------------------------------------------------------------------
Features:
  1. E-Way Bill & GSTR-2B support with remembered credentials.
  2. Smart Quarterly (QRMP) logic:
     - Detects whether taxpayer is Monthly or Quarterly (QRMP) from GST Portal.
     - For Quarterly filers:
       * April & May only contain GSTR-2A; the script automatically downloads June (covers all 3 months).
       * If incomplete quarters are selected (e.g., July & August without September), warns the user:
         "July 2025 and August 2025 cannot be downloaded because GSTR-2B is generated quarterly.
          For downloading, please select September 2025."
       * Offers to automatically include the quarter-end month (September) with one click.
  3. Automatic popup dismissal (e.g., "Principal Place of Business: Remind Me Later", Aadhaar, etc.).
  4. Direct GSTR-2B JSON download from the offline download page.
  5. Automatic JSON parsing & conversion into a portal-style Excel (.xlsx) - same layout as the GST
     portal / SPEQTA workbook: Read me, ITC Available, ITC not available, ITC Rejected, B2B, B2BA,
     B2B-CDNR, B2B-CDNRA, ISD, ISDA, IMPG, IMPGSEZ, Ecomm and the (Rejected) sheets.
  6. Automatic multi-period merging into a single consolidated Excel file:
     <output>/<GSTIN>/<GSTIN>_GSTR2B_MERGED_<range>.xlsx
"""

import csv
import json
import os
import shutil
import sys
import time
import tkinter as tk
import zipfile
from datetime import datetime
from tkinter import filedialog, messagebox, ttk

import pandas as pd
from openpyxl.utils import get_column_letter
from selenium import webdriver
from selenium.common.exceptions import (NoAlertPresentException,
                                        TimeoutException,
                                        UnexpectedAlertPresentException)
from selenium.webdriver.common.by import By
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.support.ui import Select, WebDriverWait

# ----------------------------------------------------------------------------
# Constants & URLs
# ----------------------------------------------------------------------------
LOGIN_URL = "https://ewaybillgst.gov.in/login.aspx"
HOME_URL_PART = "mainmenu.aspx"
REPORT_URL = "https://ewaybillgst.gov.in/Reports/CommomReport_month.aspx"

GST_LOGIN_URL = "https://services.gst.gov.in/services/login"
GST_RETURNS_URL = "https://return.gst.gov.in/returns/auth/dashboard"

SERVICE_EWB = "E-Way Bill (monthly)"
SERVICE_2B = "GSTR-2B"

FREQ_Q = "Quarterly Client (QRMP)"
FREQ_M = "Monthly Client"

MONTHS = ["January", "February", "March", "April", "May", "June", "July",
          "August", "September", "October", "November", "December"]

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
SAVED_FILE = os.path.join(BASE_DIR, "ewb_saved_users.json")
GST_SAVED_FILE = os.path.join(BASE_DIR, "gst_saved_users.json")

LOGIN_WAIT_SECONDS = 300      # User captcha wait
DOWNLOAD_WAIT_SECONDS = 30    # Download wait timeout


# ----------------------------------------------------------------------------
# Saved users management
# ----------------------------------------------------------------------------
def saved_path(service):
    return GST_SAVED_FILE if service == SERVICE_2B else SAVED_FILE


def load_saved(service=SERVICE_EWB):
    try:
        with open(saved_path(service), "r", encoding="utf-8") as f:
            saved = json.load(f)
        # Migrate older helper files that may contain a plaintext password.
        if isinstance(saved, dict):
            usernames_only = {username: "" for username in saved}
            if usernames_only != saved:
                with open(saved_path(service), "w", encoding="utf-8") as f:
                    json.dump(usernames_only, f, indent=2)
            return usernames_only
        return {}
    except Exception:
        return {}


def save_user(username, password, service=SERVICE_EWB):
    # Keep only the username; GST portal passwords must not be stored in plaintext.
    data = load_saved(service)
    data[username] = ""
    try:
        with open(saved_path(service), "w", encoding="utf-8") as f:
            json.dump(data, f, indent=2)
    except Exception as e:
        print(f"Warning: could not save username: {e}")


# ----------------------------------------------------------------------------
# Quarter & Period Helpers
# ----------------------------------------------------------------------------
def get_quarter_info(month_name, year_str):
    """
    Returns (quarter_num, (q_end_month, q_end_year), q_months, fy_str)
    GST Quarters (FY April to March):
      Q1: April, May, June        -> Quarter-end: June
      Q2: July, August, September  -> Quarter-end: September
      Q3: October, November, Dec   -> Quarter-end: December
      Q4: January, February, March -> Quarter-end: March
    """
    m_no = MONTHS.index(month_name) + 1
    y = int(year_str)
    if m_no in [4, 5, 6]:
        return 1, ("June", str(y)), [("April", str(y)), ("May", str(y)), ("June", str(y))], f"{y}-{str(y+1)[2:]}"
    elif m_no in [7, 8, 9]:
        return 2, ("September", str(y)), [("July", str(y)), ("August", str(y)), ("September", str(y))], f"{y}-{str(y+1)[2:]}"
    elif m_no in [10, 11, 12]:
        return 3, ("December", str(y)), [("October", str(y)), ("November", str(y)), ("December", str(y))], f"{y}-{str(y+1)[2:]}"
    else:  # 1, 2, 3
        return 4, ("March", str(y)), [("January", str(y)), ("February", str(y)), ("March", str(y))], f"{y-1}-{str(y)[2:]}"


def resolve_quarterly_periods(selected_months):
    """
    Analyzes selected months for a quarterly (QRMP) taxpayer.
    Returns:
      to_download: list of (q_end_month, q_end_year) periods to download
      quarters_info: dict mapping q_end -> dict of quarter details
      incomplete_warnings: list of dicts for quarters where months were selected but q_end was not
      skipped_months: list of (month, year, q_end) representing M1 and M2 covered by Q-end
    """
    quarters = {}
    for m, y in selected_months:
        q_num, q_end, q_months, fy_str = get_quarter_info(m, y)
        if q_end not in quarters:
            quarters[q_end] = {
                "q_num": q_num,
                "q_months": q_months,
                "fy_str": fy_str,
                "selected": []
            }
        quarters[q_end]["selected"].append((m, y))

    selected_set = set(selected_months)
    to_download = []
    skipped_months = []
    incomplete_warnings = []

    sorted_quarters = sorted(quarters.items(), key=lambda x: (int(x[0][1]), MONTHS.index(x[0][0])))
    for q_end, info in sorted_quarters:
        if q_end in selected_set:
            to_download.append(q_end)
            for m, y in info["selected"]:
                if (m, y) != q_end:
                    skipped_months.append((m, y, q_end))
        else:
            incomplete_warnings.append({
                "q_num": info["q_num"],
                "fy_str": info["fy_str"],
                "q_end": q_end,
                "selected": info["selected"],
                "missing_q_end": q_end
            })

    return to_download, quarters, incomplete_warnings, skipped_months


# ----------------------------------------------------------------------------
# Form UI
# ----------------------------------------------------------------------------
def get_inputs():
    saved = load_saved(SERVICE_EWB)
    now = datetime.now()
    result = {}

    root = tk.Tk()
    root.title("E-Way Bill / GSTR-2B Downloader")
    root.resizable(False, False)
    pad = {"padx": 8, "pady": 5}

    # Row 0: Service
    ttk.Label(root, text="Download").grid(row=0, column=0, sticky="w", **pad)
    service_var = tk.StringVar(value=SERVICE_EWB)
    service_cb = ttk.Combobox(root, textvariable=service_var, values=[SERVICE_EWB, SERVICE_2B],
                              state="readonly", width=26)
    service_cb.grid(row=0, column=1, columnspan=3, sticky="w", **pad)

    # Row 1: Username
    ttk.Label(root, text="Username").grid(row=1, column=0, sticky="w", **pad)
    user_var = tk.StringVar()
    user_cb = ttk.Combobox(root, textvariable=user_var, values=list(saved.keys()), width=28)
    user_cb.grid(row=1, column=1, columnspan=3, **pad)

    # Row 2: Password
    ttk.Label(root, text="Password").grid(row=2, column=0, sticky="w", **pad)
    pass_var = tk.StringVar()
    pass_entry = ttk.Entry(root, textvariable=pass_var, show="*", width=31)
    pass_entry.grid(row=2, column=1, columnspan=3, **pad)
    show_var = tk.BooleanVar(value=False)
    ttk.Checkbutton(root, text="Show", variable=show_var,
                    command=lambda: pass_entry.config(show="" if show_var.get() else "*"))\
        .grid(row=2, column=4, **pad)

    def on_user_pick(_e=None):
        if user_var.get() in saved:
            pass_var.set(saved[user_var.get()])

    def on_focus_out(_e=None):
        if user_var.get() in saved and not pass_var.get():
            pass_var.set(saved[user_var.get()])
    user_cb.bind("<<ComboboxSelected>>", on_user_pick)
    user_cb.bind("<FocusOut>", on_focus_out)

    # Row 3: Remember
    remember_var = tk.BooleanVar(value=True)
    ttk.Checkbutton(root, text="Remember username only (password is never saved)", variable=remember_var)\
        .grid(row=3, column=1, columnspan=3, sticky="w", **pad)

    # Row 4: Type / Filing Frequency
    type_label = ttk.Label(root, text="Type")
    type_label.grid(row=4, column=0, sticky="w", **pad)
    type_var = tk.StringVar(value="Outward")
    type_cb = ttk.Combobox(root, textvariable=type_var, values=["Outward", "Inward", "Both"],
                           state="readonly", width=12)
    type_cb.grid(row=4, column=1, sticky="w", **pad)

    freq_label = ttk.Label(root, text="Client Type")
    freq_var = tk.StringVar(value=FREQ_Q)
    freq_cb = ttk.Combobox(root, textvariable=freq_var,
                           values=[FREQ_Q, FREQ_M],
                           state="readonly", width=24)

    years = [str(y) for y in range(now.year - 6, now.year + 2)]

    # Row 5: Financial Year quick-select (April to March)
    ttk.Label(root, text="Financial Year").grid(row=5, column=0, sticky="w", **pad)
    prev_fy = now.year - 1 if now.month >= 4 else now.year - 2
    fy_values = [f"{y}-{str(y + 1)[2:]}" for y in range(now.year, now.year - 7, -1)
                 if y <= (now.year if now.month >= 4 else now.year - 1)]
    fy_var = tk.StringVar(value=f"{prev_fy}-{str(prev_fy + 1)[2:]}")
    fy_cb = ttk.Combobox(root, textvariable=fy_var, values=fy_values, state="readonly", width=12)
    fy_cb.grid(row=5, column=1, sticky="w", **pad)

    # Row 6: From
    ttk.Label(root, text="From").grid(row=6, column=0, sticky="w", **pad)
    fm = tk.StringVar(value="April")
    fy = tk.StringVar(value=str(prev_fy))
    from_m_cb = ttk.Combobox(root, textvariable=fm, values=MONTHS, state="readonly", width=12)
    from_m_cb.grid(row=6, column=1, **pad)
    from_y_cb = ttk.Combobox(root, textvariable=fy, values=years, state="readonly", width=8)
    from_y_cb.grid(row=6, column=2, **pad)

    # Row 7: To
    ttk.Label(root, text="To").grid(row=7, column=0, sticky="w", **pad)
    tm = tk.StringVar(value="March")
    ty = tk.StringVar(value=str(prev_fy + 1))
    to_m_cb = ttk.Combobox(root, textvariable=tm, values=MONTHS, state="readonly", width=12)
    to_m_cb.grid(row=7, column=1, **pad)
    to_y_cb = ttk.Combobox(root, textvariable=ty, values=years, state="readonly", width=8)
    to_y_cb.grid(row=7, column=2, **pad)

    def on_fy(_e=None):
        y = int(fy_var.get().split("-")[0])
        fm.set("April")
        fy.set(str(y))
        tm.set("March")
        ty.set(str(y + 1))
    fy_cb.bind("<<ComboboxSelected>>", on_fy)

    # Row 8: Save to
    ttk.Label(root, text="Save to").grid(row=8, column=0, sticky="w", **pad)
    out_var = tk.StringVar(value=os.path.join(os.path.expanduser("~"), "Desktop", "GST_Downloads"))
    ttk.Entry(root, textvariable=out_var, width=31).grid(row=8, column=1, columnspan=2, **pad)
    ttk.Button(root, text="Browse",
               command=lambda: out_var.set(filedialog.askdirectory() or out_var.get()))\
        .grid(row=8, column=3, **pad)

    def on_service(_e=None):
        saved.clear()
        svc = service_var.get()
        saved.update(load_saved(svc))
        user_cb["values"] = list(saved.keys())
        user_var.set("")
        pass_var.set("")
        if svc == SERVICE_2B:
            type_label.grid_remove()
            type_cb.grid_remove()
            freq_label.grid(row=4, column=0, sticky="w", **pad)
            freq_cb.grid(row=4, column=1, columnspan=2, sticky="w", **pad)
            out_var.set(os.path.join(os.path.expanduser("~"), "Desktop", "GSTR2B_Downloads"))
        else:
            freq_label.grid_remove()
            freq_cb.grid_remove()
            type_label.grid(row=4, column=0, sticky="w", **pad)
            type_cb.grid(row=4, column=1, sticky="w", **pad)
            type_cb.config(state="readonly")
            out_var.set(os.path.join(os.path.expanduser("~"), "Desktop", "EWB_Downloads"))
    service_cb.bind("<<ComboboxSelected>>", on_service)

    def start():
        u = user_var.get().strip()
        if not u:
            messagebox.showerror("Missing", "Enter username")
            return
        start_idx = int(fy.get()) * 12 + MONTHS.index(fm.get())
        end_idx = int(ty.get()) * 12 + MONTHS.index(tm.get())
        if start_idx > end_idx:
            messagebox.showerror("Range", "'From' is after 'To'")
            return

        months_list = [(MONTHS[i % 12], str(i // 12)) for i in range(start_idx, end_idx + 1)]
        service = service_var.get()
        frequency = freq_var.get()

        # Check quarterly warnings if Quarterly is explicitly selected
        if service == SERVICE_2B and frequency == FREQ_Q:
            to_dl, qinfo, warnings, skipped = resolve_quarterly_periods(months_list)
            if warnings:
                w_lines = []
                for w in warnings:
                    sel_str = " and ".join([f"{m} {y}" for m, y in w["selected"]])
                    q_end_str = f"{w['missing_q_end'][0]} {w['missing_q_end'][1]}"
                    w_lines.append(f"• {sel_str} cannot be downloaded because GSTR-2B is generated quarterly.\n"
                                   f"  To download data for this quarter, please select {q_end_str}.")
                warning_text = "\n\n".join(w_lines)
                opt = messagebox.askyesno(
                    "Quarterly GSTR-2B Warning",
                    f"{warning_text}\n\n"
                    f"Would you like to automatically extend the 'To' period to include the missing quarter-end month(s)?"
                )
                if opt:
                    last_missing = warnings[-1]["missing_q_end"]
                    tm.set(last_missing[0])
                    ty.set(last_missing[1])
                    end_idx = int(ty.get()) * 12 + MONTHS.index(tm.get())
                    months_list = [(MONTHS[i % 12], str(i // 12)) for i in range(start_idx, end_idx + 1)]
                else:
                    messagebox.showinfo("Notice", "The incomplete quarter months will be skipped during download.")

        types = ["GSTR-2B"] if service == SERVICE_2B else (
            ["Outward", "Inward"] if type_var.get() == "Both" else [type_var.get()]
        )

        ok = messagebox.askyesno(
            "Confirm Download",
            f"Service: {service}\n"
            f"User: {u}\n"
            f"Period: {fm.get()} {fy.get()} to {tm.get()} {ty.get()} ({len(months_list)} months)\n"
            f"{f'Filing Mode: {frequency}' if service == SERVICE_2B else f'Type: {type_var.get()}'}\n\n"
            f"Start Download?"
        )
        if not ok:
            return

        result.update(
            service=service, username=u, password=pass_var.get().strip(),
            remember=remember_var.get(),
            types=types,
            frequency=frequency,
            months=months_list,
            out=out_var.get().strip(), manual_pw=manual_var.get(),
        )
        root.destroy()

    manual_var = tk.BooleanVar(value=False)
    ttk.Checkbutton(root, text="I will type the password myself (script fills username only)",
                    variable=manual_var).grid(row=9, column=0, columnspan=4, sticky="w", **pad)
    ttk.Button(root, text="Start Download", command=start).grid(row=10, column=1, columnspan=2, pady=12)
    root.mainloop()
    return result or None


# ----------------------------------------------------------------------------
# Browser helpers & Automation
# ----------------------------------------------------------------------------
def make_driver(temp_dl_dir):
    from selenium.webdriver.chrome.service import Service

    def build_opts():
        opts = webdriver.ChromeOptions()
        opts.add_experimental_option("prefs", {
            "download.default_directory": temp_dl_dir,
            "download.prompt_for_download": False,
            "download.directory_upgrade": True,
            "safebrowsing.enabled": True,
            "profile.default_content_setting_values.automatic_downloads": 1,
        })
        opts.add_argument("--start-maximized")
        opts.add_experimental_option("excludeSwitches", ["enable-logging"])
        return opts

    errors = []
    try:
        print("Starting Chrome (Selenium Manager)...")
        return webdriver.Chrome(options=build_opts())
    except Exception as e:
        errors.append(f"Method 1: {str(e)[:300]}")

    try:
        print("Starting Chrome (webdriver-manager)...")
        from webdriver_manager.chrome import ChromeDriverManager
        return webdriver.Chrome(service=Service(ChromeDriverManager().install()),
                                options=build_opts())
    except Exception as e:
        errors.append(f"Method 2: {str(e)[:300]}")

    local = os.path.join(BASE_DIR, "chromedriver.exe")
    if os.path.exists(local):
        try:
            print("Starting Chrome (local chromedriver.exe)...")
            return webdriver.Chrome(service=Service(local), options=build_opts())
        except Exception as e:
            errors.append(f"Method 3: {str(e)[:300]}")
    else:
        errors.append("Method 3: chromedriver.exe not found next to script")

    raise RuntimeError("Could not start Chrome.\n\n" + "\n\n".join(errors))


def first_present(driver, selectors):
    for by, sel in selectors:
        els = driver.find_elements(by, sel)
        if els:
            return els[0]
    return None


def dismiss_alert(driver):
    try:
        a = driver.switch_to.alert
        text = a.text
        a.accept()
        return text
    except NoAlertPresentException:
        return None


def dismiss_gst_popups(driver):
    """
    Dismisses modal popups that appear after login on GST portal:
      - 'GST System is collecting metadata for the Principal Place of Business: NO-REMIND ME LATER'
      - Aadhaar authentication alerts
      - General informational popups
    """
    popup_xps = [
        "//button[contains(translate(normalize-space(.), 'abcdefghijklmnopqrstuvwxyz', 'ABCDEFGHIJKLMNOPQRSTUVWXYZ'), 'REMIND ME LATER')]",
        "//button[contains(translate(normalize-space(.), 'abcdefghijklmnopqrstuvwxyz', 'ABCDEFGHIJKLMNOPQRSTUVWXYZ'), 'NO-REMIND ME LATER')]",
        "//button[contains(translate(normalize-space(.), 'abcdefghijklmnopqrstuvwxyz', 'ABCDEFGHIJKLMNOPQRSTUVWXYZ'), 'LATER')]",
        "//button[contains(translate(normalize-space(.), 'abcdefghijklmnopqrstuvwxyz', 'ABCDEFGHIJKLMNOPQRSTUVWXYZ'), 'CLOSE')]",
        "//button[contains(@class, 'close') or @id='popup_close']",
        "//a[contains(translate(normalize-space(.), 'abcdefghijklmnopqrstuvwxyz', 'ABCDEFGHIJKLMNOPQRSTUVWXYZ'), 'REMIND ME LATER')]"
    ]
    for xp in popup_xps:
        try:
            btns = [b for b in driver.find_elements(By.XPATH, xp) if b.is_displayed()]
            if btns:
                driver.execute_script("arguments[0].click();", btns[0])
                time.sleep(1)
        except Exception:
            pass


def type_slow(el, text):
    el.click()
    el.clear()
    for ch in text:
        el.send_keys(ch)
        time.sleep(0.06)


def js_set(driver, el, text):
    driver.execute_script("""
        var el = arguments[0], v = arguments[1];
        var setter = Object.getOwnPropertyDescriptor(window.HTMLInputElement.prototype, 'value').set;
        setter.call(el, v);
        ['input', 'keyup', 'change'].forEach(function (t) {
            el.dispatchEvent(new Event(t, {bubbles: true}));
        });
    """, el, text)


def find_login_boxes(driver):
    user_box = first_present(driver, [
        (By.ID, "txt_username"), (By.NAME, "txt_username"),
        (By.CSS_SELECTOR, "input[type='text']"),
    ])
    pass_box = first_present(driver, [
        (By.ID, "txt_password"), (By.NAME, "txt_password"),
        (By.CSS_SELECTOR, "input[type='password']"),
    ])
    return user_box, pass_box


def gst_login_boxes(driver):
    user_box = first_present(driver, [
        (By.ID, "username"), (By.NAME, "username"), (By.CSS_SELECTOR, "input[type='text']")])
    pass_box = first_present(driver, [
        (By.ID, "user_pass"), (By.NAME, "user_pass"), (By.CSS_SELECTOR, "input[type='password']")])
    return user_box, pass_box


def fill_credentials(driver, username, password, fill_pw=True, boxes_fn=None):
    username = username.strip()
    password = password.strip()
    find = boxes_fn or find_login_boxes

    methods = [
        ("typing", type_slow),
        ("direct value set", lambda el, t: js_set(driver, el, t)),
    ]
    ok = False
    for name, fn in methods:
        user_box, pass_box = find(driver)
        if user_box:
            fn(user_box, username)
        if fill_pw and pass_box:
            fn(pass_box, password)
        u_val = (user_box.get_attribute("value") if user_box else "") or ""
        p_val = (pass_box.get_attribute("value") if pass_box else "") or ""
        if u_val == username and (not fill_pw or p_val == password):
            ok = True
            break
        time.sleep(1)

    if not fill_pw:
        print(">>> Type your PASSWORD and the CAPTCHA, then click Login.")
        pb = find(driver)[1]
        if pb:
            pb.click()
        return

    cap = first_present(driver, [(By.ID, "txtCaptcha"), (By.NAME, "txtCaptcha"),
                                 (By.ID, "captcha"), (By.NAME, "captcha")])
    if cap:
        cap.click()


def page_ready(driver):
    WebDriverWait(driver, 30).until(
        lambda d: d.find_elements(By.CSS_SELECTOR, "input[type='password']"))
    WebDriverWait(driver, 30).until(
        lambda d: d.execute_script("return document.readyState") == "complete")


def fill_login(driver, username, password, fill_pw=True):
    driver.get(LOGIN_URL)
    page_ready(driver)
    time.sleep(1.5)
    driver.refresh()
    page_ready(driver)
    time.sleep(1.5)
    fill_credentials(driver, username, password, fill_pw)


def gst_login(driver, username, password, fill_pw=True):
    driver.get(GST_LOGIN_URL)
    page_ready(driver)
    time.sleep(2)
    fill_credentials(driver, username, password, fill_pw, boxes_fn=gst_login_boxes)


def wait_for_login(driver, refill=None, home_part=HOME_URL_PART):
    print("\n>>> Enter the password if needed, CAPTCHA, and OTP if asked, then click Login. Waiting...")
    end = time.time() + LOGIN_WAIT_SECONDS
    refills = 0
    while time.time() < end:
        try:
            url = driver.current_url.lower()
            if home_part in url:
                time.sleep(1.5)
                return True
            if refill and refills < 5 and "login" in url:
                pb = first_present(driver, [
                    (By.ID, "txt_password"), (By.NAME, "txt_password"),
                    (By.ID, "user_pass"), (By.NAME, "user_pass"),
                    (By.CSS_SELECTOR, "input[type='password']"),
                ])
                if pb is not None and (pb.get_attribute("value") or "") == "":
                    time.sleep(1.5)
                    refills += 1
                    print("   Login page reloaded. Refilling credentials - please enter new CAPTCHA.")
                    refill()
        except UnexpectedAlertPresentException:
            dismiss_alert(driver)
        except Exception:
            pass
        time.sleep(1)
    return False


def get_gstin(driver, fallback):
    try:
        text = driver.find_element(By.XPATH, "//*[contains(text(),'GSTIN')]").text
        part = text.split("GSTIN")[1].replace(":", " ").split()[0]
        if len(part) == 15:
            return part
    except Exception:
        pass
    return fallback


def list_files(folder):
    if not os.path.exists(folder):
        return set()
    return set(os.listdir(folder))


def wait_new_file(folder, before, timeout):
    end = time.time() + timeout
    while time.time() < end:
        new = [f for f in (list_files(folder) - before)
               if not f.lower().endswith((".crdownload", ".tmp"))]
        if new:
            time.sleep(1.0)
            return new[0]
        time.sleep(0.5)
    return None


def click_xpath(driver, xps, timeout=15):
    end = time.time() + timeout
    while True:
        for xp in xps:
            try:
                els = [e for e in driver.find_elements(By.XPATH, xp) if e.is_displayed()]
            except Exception:
                els = []
            if els:
                try:
                    driver.execute_script("arguments[0].scrollIntoView({block:'center'});", els[0])
                    time.sleep(0.3)
                    driver.execute_script("arguments[0].click();", els[0])
                    return True
                except Exception:
                    pass
        if time.time() >= end:
            return False
        time.sleep(0.5)


def gst_select(driver, names, index, text, starts=False):
    def find():
        for n in names:
            els = driver.find_elements(By.NAME, n)
            if els:
                return els[0]
        sels = [e for e in driver.find_elements(By.TAG_NAME, "select") if e.is_displayed()]
        return sels[index] if len(sels) > index else None

    el = find()
    if el is None:
        raise RuntimeError(f"Dropdown '{names[0]}' not found")
    sel = Select(el)
    target = None
    for o in sel.options:
        t = o.text.strip()
        if (t.startswith(text) if starts else t == text):
            target = o
            break
    if target is None:
        have = [o.text.strip() for o in sel.options][:8]
        raise RuntimeError(f"Option '{text}' not in dropdown '{names[0]}' (Available: {have})")
    if sel.first_selected_option.text.strip() == target.text.strip():
        return
    val = target.get_attribute("value")
    if val is not None and val != "":
        sel.select_by_value(val)
    else:
        target.click()
    time.sleep(1.5)


from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

STATE_CODES = {
    "01": "Jammu and Kashmir", "02": "Himachal Pradesh", "03": "Punjab", "04": "Chandigarh",
    "05": "Uttarakhand", "06": "Haryana", "07": "Delhi", "08": "Rajasthan", "09": "Uttar Pradesh",
    "10": "Bihar", "11": "Sikkim", "12": "Arunachal Pradesh", "13": "Nagaland", "14": "Manipur",
    "15": "Mizoram", "16": "Tripura", "17": "Meghalaya", "18": "Assam", "19": "West Bengal",
    "20": "Jharkhand", "21": "Odisha", "22": "Chhattisgarh", "23": "Madhya Pradesh", "24": "Gujarat",
    "25": "Daman and Diu", "26": "Dadra and Nagar Haveli and Daman and Diu", "27": "Maharashtra",
    "28": "Andhra Pradesh (Old)", "29": "Karnataka", "30": "Goa", "31": "Lakshadweep", "32": "Kerala",
    "33": "Tamil Nadu", "34": "Puducherry", "35": "Andaman and Nicobar Islands", "36": "Telangana",
    "37": "Andhra Pradesh", "38": "Ladakh", "96": "Foreign Country", "97": "Other Territory",
}
INV_TYPES = {"R": "Regular", "SEWP": "SEZ supplies with payment", "SEWOP": "SEZ supplies without payment",
             "DE": "Deemed Exports", "CBW": "Intra-State supplies attracting IGST"}
NOTE_TYPES = {"C": "Credit Note", "D": "Debit Note"}
ITC_AVL = {"Y": "Yes", "N": "No", "T": "Temporarily Not Available"}
REASONS = {"P": "POS and supplier state are same but recipient state is different",
           "C": "Return filed post annual return cut-off date"}
STD_RATES = [0, 0.1, 0.25, 1, 1.5, 3, 5, 6, 7.5, 12, 18, 28, 40]
MON = ["", "January", "February", "March", "April", "May", "June", "July", "August",
       "September", "October", "November", "December"]
MON3 = ["", "Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]

NAVY, YELLOW, ORANGE, DKBLUE, DKORANGE = "203764", "FFF2CC", "F4B084", "1F4E78", "C65911"
THIN = Side(style="thin", color="000000")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
NUM = "[$-en-US] #,##0.00"


def _fill(c):
    return PatternFill("solid", fgColor=c)


def _dmy(s):
    return (s or "").replace("-", "/")


def _prd_long(p):            # "022026" -> "February'26"
    p = p or ""
    return f"{MON[int(p[:2])]}'{p[-2:]}" if len(p) == 6 and p[:2].isdigit() else p


def _prd_short(p):           # "032026" -> "Mar- 26"
    p = p or ""
    return f"{MON3[int(p[:2])]}- {p[-2:]}" if len(p) == 6 and p[:2].isdigit() else p


def _pos(c):
    c = str(c or "").zfill(2) if str(c or "").isdigit() else str(c or "")
    return STATE_CODES.get(c, c)


def _yn(v):
    return "Yes" if str(v).upper() == "Y" else ("No" if str(v).upper() == "N" else "")


def _rate(d):
    """Rate (%) - from items if present, else derived from tax/taxable value; 'Multi Rate' otherwise."""
    items = d.get("items") or []
    rts = {str(i.get("rt")) for i in items if i.get("rt") is not None}
    if len(rts) == 1:
        r = float(next(iter(rts)))
        return str(int(r)) if r == int(r) else str(r)
    if len(rts) > 1:
        return "Multi Rate"
    tx = float(d.get("txval") or 0)
    tax = sum(float(d.get(k) or 0) for k in ("igst", "cgst", "sgst"))
    if tx <= 0:
        return "" if tax == 0 else "Multi Rate"
    for r in STD_RATES:
        if abs(tax - tx * r / 100) <= 1.0:
            return str(int(r)) if r == int(r) else str(r)
    return "Multi Rate"


def _g(d, k, default=0):
    v = d.get(k)
    return default if v is None else v


def _title(ws, sub, ncols, hdr_rows):
    ws.merge_cells(start_row=1, start_column=1, end_row=3, end_column=ncols)
    c = ws.cell(1, 1, "Goods and Services Tax  - GSTR-2B")
    c.font = Font(name="Arial", size=22, color="FFFFFF")
    c.fill = _fill(NAVY)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.merge_cells(start_row=4, start_column=1, end_row=4, end_column=ncols)
    c = ws.cell(4, 1, sub)
    c.font = Font(name="Arial", size=11, bold=True)
    c.fill = _fill(YELLOW)
    c.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
    for r in range(1, 5):
        for col in range(1, ncols + 1):
            ws.cell(r, col).border = BORDER
            if r < 4:
                ws.cell(r, col).fill = _fill(NAVY)
            else:
                ws.cell(r, col).fill = _fill(YELLOW)


def _header(ws, spec, ncols, nrows):
    """spec: list of (r1, c1, r2, c2, text). Draws navy header cells with merges."""
    for r in range(5, 5 + nrows):
        for col in range(1, ncols + 1):
            c = ws.cell(r, col)
            c.fill = _fill(NAVY)
            c.border = BORDER
            c.font = Font(name="Arial", size=9, bold=True, color="FFFFFF")
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r1, c1, r2, c2, text in spec:
        if (r1, c1) != (r2, c2):
            ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
        ws.cell(r1, c1, text)


def _sheet(wb, name, sub, ncols, hdr_spec, hdr_rows, widths=None):
    ws = wb.create_sheet(name)
    _title(ws, sub, ncols, hdr_rows)
    _header(ws, hdr_spec, ncols, hdr_rows)
    ws.freeze_panes = ws.cell(5 + hdr_rows, 1)
    for i in range(1, ncols + 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(i, 16)
    return ws


def _put(ws, row, values, num_cols=(), text_cols=()):
    for i, v in enumerate(values, 1):
        c = ws.cell(row, i, v)
        c.font = Font(name="Calibri", size=11)
        if i in num_cols:
            c.number_format = NUM
            c.alignment = Alignment(horizontal="right", vertical="bottom")
        elif i in text_cols:
            c.alignment = Alignment(horizontal="left")
        else:
            c.alignment = Alignment(horizontal="right")


# ---------------------------------------------------------------- summary sheets
def _tot(sec, key):
    return [_g(sec.get(key, {}) or {}, k) for k in ("igst", "cgst", "sgst", "cess")] if isinstance(sec.get(key), dict) else [0, 0, 0, 0]


def _sum_row_vals(node):
    node = node if isinstance(node, dict) else {}
    return [_g(node, "igst"), _g(node, "cgst"), _g(node, "sgst"), _g(node, "cess")]


def _summary_sheet(wb, name, sub_title, banner, part_defs):
    """part_defs: list of dict(kind='part'|'main'|'detail', ...) - see callers."""
    ws = wb.create_sheet(name)
    for col, w in zip("ABCDEFGHIJ", (6.14, 29, 9, 17, 17, 17, 14, 14, 14, 14)):
        ws.column_dimensions[col].width = w
    ws.merge_cells("A1:J1")
    c = ws.cell(1, 1, "FORM GSTR-2B")
    c.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    c.fill = _fill("002060")
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.merge_cells("A2:J4")
    c = ws.cell(2, 1, banner)
    c.font = Font(name="Arial", size=10, bold=True)
    c.fill = _fill(ORANGE)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.merge_cells("A5:J5")
    c = ws.cell(5, 1, sub_title)
    c.font = Font(name="Calibri", size=10, color="FFFFFF")
    c.fill = _fill(DKBLUE)
    c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    heads = ["S.no.", "Heading", "GSTR-3B table", "Integrated Tax  (₹)", "Central Tax (₹)",
             "State/UT Tax (₹)", "Cess  (₹)", "Advisory"]
    for i, h in enumerate(heads, 1):
        c = ws.cell(6, i, h)
        c.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        c.fill = _fill("002060")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.merge_cells("H6:J6")
    ws.row_dimensions[6].height = 33.75
    r = 7
    for p in part_defs:
        k = p["kind"]
        if k == "banner":
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=10)
            c = ws.cell(r, 1, p["text"])
            c.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            c.fill = _fill(DKORANGE)
            c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            r += 1
        elif k == "part":
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=10)
            ws.cell(r, 1, p["no"])
            ws.cell(r, 2, p["text"])
            for col in (1, 2):
                c = ws.cell(r, col)
                c.font = Font(name="Arial", size=10, bold=True)
                c.fill = _fill(ORANGE)
                c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            ws.row_dimensions[r].height = 25.5
            r += 1
        elif k == "main":
            ws.cell(r, 1, p["no"])
            ws.cell(r, 2, p["text"]).font = Font(name="Calibri", size=10, bold=True)
            ws.cell(r, 3, p["tbl"]).font = Font(name="Calibri", size=10, bold=True)
            for j, v in enumerate(p["vals"]):
                ws.cell(r, 4 + j, v)
            ws.cell(r, 8, p["adv"])
            ws.merge_cells(start_row=r, start_column=8, end_row=r, end_column=10)
            ws.cell(r, 1).font = Font(name="Calibri", size=10)
            ws.cell(r, 3).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.cell(r, 1).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.cell(r, 2).alignment = Alignment(wrap_text=True)
            ws.cell(r, 8).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.cell(r, 8).font = Font(name="Calibri", size=10)
            ws.row_dimensions[r].height = 36
            first_detail = r + 1
            r += 1
            for lbl, vals in p["details"]:
                ws.cell(r, 2, lbl).alignment = Alignment(wrap_text=True)
                ws.cell(r, 2).font = Font(name="Calibri", size=10)
                for j, v in enumerate(vals):
                    ws.cell(r, 4 + j, v)
                r += 1
            last_detail = r - 1
            ws.merge_cells(start_row=first_detail, start_column=1, end_row=last_detail, end_column=1)
            ws.cell(first_detail, 1, "Details").alignment = Alignment(horizontal="center", vertical="center")
            ws.cell(first_detail, 1).font = Font(name="Calibri", size=10)
            ws.merge_cells(start_row=first_detail, start_column=8, end_row=last_detail, end_column=10)
        for rr in range(r - 1 if k == "banner" or k == "part" else (first_detail - 1), r):
            pass
    # number format + borders for the whole table
    for row in ws.iter_rows(min_row=1, max_row=r - 1, min_col=1, max_col=10):
        for c in row:
            c.border = BORDER
            if 4 <= c.column <= 7 and c.row >= 7 and isinstance(c.value, (int, float)):
                c.number_format = NUM
                c.alignment = Alignment(horizontal="right", vertical="bottom", wrap_text=True)
    return ws


def _build_summaries(wb, sm):
    sm = sm or {}
    av = sm.get("itcavl", {}) or {}
    un = sm.get("itcunavl", {}) or {}
    rj = sm.get("itcrej", {}) or {}

    def sec(root, key):
        return root.get(key, {}) if isinstance(root.get(key), dict) else {}

    def det(root_sec, *keys):
        for k in keys:
            if isinstance(root_sec.get(k), dict):
                return _sum_row_vals(root_sec[k])
        return [0, 0, 0, 0]

    # ---- ITC Available
    nr, rv, isd, imp, oth = (sec(av, "nonrevsup"), sec(av, "revsup"), sec(av, "isdsup"),
                             sec(av, "impsup"), sec(av, "othersup"))
    parts = [
        {"kind": "banner", "text": "Credit which may be availed under FORM GSTR-3B"},
        {"kind": "part", "no": "Part A", "text": "ITC Available - Credit may be claimed in relevant headings in GSTR-3B"},
        {"kind": "main", "no": "I", "text": "All other ITC - Supplies from registered persons other than reverse charge",
         "tbl": "4(A)(5)", "vals": _sum_row_vals(nr),
         "adv": "Net input tax credit may be availed under Table 4(A)(5) of FORM GSTR-3B.",
         "details": [("B2B - Invoices", det(nr, "b2b")), ("B2B - Debit notes", det(nr, "cdnr")),
                     ("Ecomm", det(nr, "ecom")), ("B2B - Invoices (Amendment)", det(nr, "b2ba")),
                     ("B2B - Debit notes (Amendment)", det(nr, "cdnra")), ("Ecomm (Amendment)", det(nr, "ecoma"))]},
        {"kind": "main", "no": "II", "text": "Inward Supplies from ISD", "tbl": "4(A)(4)", "vals": _sum_row_vals(isd),
         "adv": "Net input tax credit may be availed under Table 4(A)(4) of FORM GSTR-3B.",
         "details": [("ISD - Invoices", det(isd, "isd")), ("ISD - Invoices (Amendment)", det(isd, "isda"))]},
        {"kind": "main", "no": "III", "text": "Inward Supplies liable for reverse charge", "tbl": "3.1(d) \n 4(A)(3)",
         "vals": _sum_row_vals(rv),
         "adv": "These supplies shall be declared in Table 3.1(d) of FORM GSTR-3B for payment of tax. Net input tax credit may be availed under Table 4(A)(3) of FORM GSTR-3B on payment of tax.",
         "details": [("B2B - Invoices", det(rv, "b2b")), ("B2B - Debit notes", det(rv, "cdnr")),
                     ("B2B - Invoices (Amendment)", det(rv, "b2ba")), ("B2B - Debit notes (Amendment)", det(rv, "cdnra"))]},
        {"kind": "main", "no": "IV", "text": "Import of Goods", "tbl": "4(A)(1)", "vals": _sum_row_vals(imp),
         "adv": "Net input tax credit may be availed under Table 4(A)(1) of FORM GSTR-3B.",
         "details": [("IMPG - Import of goods from overseas", det(imp, "impg")), ("IMPG (Amendment)", det(imp, "impga")),
                     ("IMPGSEZ - Import of goods from SEZ", det(imp, "impgsez")), ("IMPGSEZ (Amendment)", det(imp, "impgseza"))]},
        {"kind": "part", "no": "Part B", "text": "ITC Reversal - Credit should be reversed in relevant headings in GSTR-3B"},
        {"kind": "main", "no": "I", "text": "Others", "tbl": "4(B)(2)", "vals": _sum_row_vals(oth),
         "adv": "If this is positive, Credit shall be reversed under Table 4(B)(2) of FORM GSTR-3B.",
         "details": [("B2B - Credit notes", det(oth, "cdnr")), ("B2B - Credit notes (Amendment)", det(oth, "cdnra")),
                     ("B2B - Credit notes (Reverse charge)", det(oth, "cdnrrc", "cdnrrev")),
                     ("B2B - Credit notes (Reverse charge)(Amendment)", det(oth, "cdnrarc", "cdnrarev")),
                     ("ISD - Credit notes", det(oth, "isd")), ("ISD - Credit notes (Amendment)", det(oth, "isda"))]},
    ]
    banner1 = ("FORM GSTR-2B has been generated on the basis of the information furnished by your suppliers in "
               "their respective FORMS GSTR-1,5 and 6. It also contains information on imports of goods from the "
               "ICEGATE system. This information is for guidance purposes only.")
    _summary_sheet(wb, "ITC Available", "FORM SUMMARY - ITC Available", banner1, parts)

    # ---- ITC not available
    nr, rv, isd, oth = sec(un, "nonrevsup"), sec(un, "revsup"), sec(un, "isdsup"), sec(un, "othersup")
    adv_na = " Such credit shall not be taken and has to be reported in table 4(D)(2) of FORM GSTR-3B."
    parts = [
        {"kind": "banner", "text": "Credit which may not be availed under FORM GSTR-3B"},
        {"kind": "part", "no": "Part A", "text": "ITC Not Available"},
        {"kind": "main", "no": "I", "text": "All other ITC - Supplies from registered persons other than reverse charge",
         "tbl": "NA", "vals": _sum_row_vals(nr), "adv": adv_na,
         "details": [("B2B - Invoices", det(nr, "b2b")), ("B2B - Debit notes", det(nr, "cdnr")),
                     ("B2B - Invoices (Amendment)", det(nr, "b2ba")), ("B2B - Debit notes (Amendment)", det(nr, "cdnra"))]},
        {"kind": "main", "no": "II", "text": "Inward Supplies from ISD", "tbl": "NA", "vals": _sum_row_vals(isd), "adv": adv_na,
         "details": [("ISD - Invoices", det(isd, "isd")), ("ISD - Invoices (Amendment)", det(isd, "isda"))]},
        {"kind": "main", "no": "III", "text": "Inward Supplies liable for reverse charge", "tbl": "3.1(d)", "vals": _sum_row_vals(rv),
         "adv": "These supplies shall be declared in Table 3.1(d) of FORM GSTR-3B for payment of tax. \n However, credit will not be available on the same and has to be reported in table 4(D)(2) of FORM GSTR-3B.",
         "details": [("B2B - Invoices", det(rv, "b2b")), ("B2B - Debit notes", det(rv, "cdnr")),
                     ("B2B - Invoices (Amendment)", det(rv, "b2ba")), ("B2B - Debit notes (Amendment)", det(rv, "cdnra"))]},
        {"kind": "part", "no": "Part B", "text": "ITC Not Available - Credit notes should be net off against relevant ITC available headings in GSTR-3B"},
        {"kind": "main", "no": "I", "text": "Others", "tbl": "4(B)(2)", "vals": _sum_row_vals(oth),
         "adv": "Credit Notes should be net-off against relevant ITC available tables [Table 4A(3,4,5)].",
         "details": [("B2B - Credit notes", det(oth, "cdnr")), ("B2B - Credit notes (Amendment)", det(oth, "cdnra")),
                     ("B2B - Credit notes (Reverse charge)", det(oth, "cdnrrc", "cdnrrev")),
                     ("B2B - Credit notes (Reverse charge)(Amendment)", det(oth, "cdnrarc", "cdnrarev")),
                     ("ISD - Credit notes", det(oth, "isd")), ("ISD - Credit notes (Amendment)", det(oth, "isda"))]},
    ]
    _summary_sheet(wb, "ITC not available", "FORM SUMMARY - ITC Not Available", banner1, parts)

    # ---- ITC Rejected
    nr, isd, oth = sec(rj, "nonrevsup"), sec(rj, "isdsup"), sec(rj, "othersup")
    adv_rj = "Input tax credit cannot be availed in FORM GSTR-3B."
    banner2 = ("FORM GSTR-2B has been generated on the basis of the information furnished by your suppliers in "
               "their respective FORMS GSTR-1/IFF including E-Commerce supplies, GSTR-1A, 5 and 6. It also contains "
               "information on imports of goods from the ICEGATE system. This information is for guidance purposes only.")
    parts = [
        {"kind": "banner", "text": "Credit which is rejected on IMS Dashboard"},
        {"kind": "part", "no": "Part A", "text": "ITC Rejected - Others"},
        {"kind": "main", "no": "I", "text": "All other ITC - Supplies from registered persons other than reverse charge (IMS)",
         "tbl": "NA", "vals": _sum_row_vals(nr), "adv": adv_rj,
         "details": [("B2B - Invoices (IMS)", det(nr, "b2b")), ("B2B - Debit notes (IMS)", det(nr, "cdnr")),
                     ("ECO - Documents (IMS)", det(nr, "ecom")), ("B2B - Invoices (Amendment) (IMS)", det(nr, "b2ba")),
                     ("B2B - Debit notes (Amendment) (IMS)", det(nr, "cdnra")),
                     ("ECO - Documents (Amendment) (IMS)", det(nr, "ecoma"))]},
        {"kind": "main", "no": "II", "text": "Inward Supplies from ISD", "tbl": "NA", "vals": _sum_row_vals(isd), "adv": adv_rj,
         "details": [("ISD - Invoices", det(isd, "isd")), ("ISD - Invoices (Amendment)", det(isd, "isda"))]},
        {"kind": "part", "no": "Part B", "text": "Rejected Records - Credit notes rejected on IMS Dashboard"},
        {"kind": "main", "no": "I", "text": "Others", "tbl": "NA", "vals": _sum_row_vals(oth),
         "adv": "These Credit Notes are not eligible to net-off against relevant ITC available tables [Table 4A(4,5)].",
         "details": [("B2B - Credit notes (IMS)", det(oth, "cdnr")), ("B2B - Credit notes (Amendment) (IMS)", det(oth, "cdnra")),
                     ("ISD - Credit notes", det(oth, "isd")), ("ISD - Credit notes (Amendment)", det(oth, "isda"))]},
    ]
    _summary_sheet(wb, "ITC Rejected", "FORM SUMMARY - ITC Rejected", banner2, parts)


# ---------------------------------------------------------------- detail sheets
def _is_rej(doc):
    return str(doc.get("imsStatus", "")).upper() == "R"


def _inv_rows(suppliers, period, rejected=False):
    rows = []
    for s in suppliers or []:
        for i in s.get("inv", []) or []:
            if _is_rej(i) != rejected:
                continue
            rows.append((s, i))
    return rows


def _note_rows(suppliers, rejected=False):
    rows = []
    for s in suppliers or []:
        for n in s.get("nt", []) or []:
            if _is_rej(n) != rejected:
                continue
            rows.append((s, n))
    return rows


def build_portal_workbook(data, quarterly=False):
    """data = parsed JSON (the whole file or its 'data' node). Returns an openpyxl Workbook."""
    root = data.get("data", data)
    dd = root.get("docdata", {}) or {}
    period = _prd_short(root.get("rtnprd", ""))
    gstin = root.get("gstin", "")
    rtn = root.get("rtnprd", "")

    wb = Workbook()
    wb.remove(wb.active)

    # ---- Read me
    ws = wb.create_sheet("Read me")
    for col, w in zip("ABCDEF", (10.57, 25.14, 26.43, 29, 22.14, 23.71)):
        ws.column_dimensions[col].width = w
    ws.merge_cells("A1:F3")
    c = ws.cell(1, 1, "Goods and Services Tax  - GSTR-2B")
    c.font = Font(name="Arial", size=22, color="FFFFFF")
    c.fill = _fill(NAVY)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    mm = int(rtn[:2]) if len(rtn) == 6 and rtn[:2].isdigit() else 0
    yy = int(rtn[2:]) if len(rtn) == 6 else 0
    fy = f"{yy if mm >= 4 else yy - 1} - {(yy if mm >= 4 else yy - 1) + 1}" if mm else ""
    q = ((mm - 4) % 12) // 3 + 1 if mm else 0
    tax_period = f"Quarter-{q}" if quarterly else (MON[mm] if mm else rtn)
    info = [("Financial Year", fy), ("Tax Period", tax_period), ("GSTIN", gstin),
            ("Legal Name", root.get("lgnm", "")), ("Trade Name (if any)", root.get("trdnm", "")),
            ("Date of generation", _dmy(root.get("gendt", "")))]
    for n, (k, v) in enumerate(info):
        r = 4 + n
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=6)
        a = ws.cell(r, 1, k)
        a.alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
        b = ws.cell(r, 3, v)
        b.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        for col in range(1, 7):
            ws.cell(r, col).fill = _fill(YELLOW)
            ws.cell(r, col).border = BORDER
            ws.cell(r, col).font = Font(name="Calibri", size=11)
    ws.cell(11, 1, "Sheets: ITC Available | ITC not available | ITC Rejected | B2B | B2BA | B2B-CDNR | B2B-CDNRA | ISD | ISDA | IMPG | IMPGSEZ | Ecomm ... (same layout as the GST portal Excel)")

    # ---- Summary sheets
    _build_summaries(wb, root.get("itcsumm", {}))

    # ---- B2B
    hdr = [(5, 1, 6, 1, "GSTIN of supplier"), (5, 2, 6, 2, "Trade/Legal name"), (5, 3, 5, 6, "Invoice Details"),
           (6, 3, 6, 3, "Invoice number"), (6, 4, 6, 4, "Invoice type"), (6, 5, 6, 5, "Invoice Date"),
           (6, 6, 6, 6, "Invoice Value(₹)"), (5, 7, 6, 7, "Place of supply"),
           (5, 8, 6, 8, "Supply Attract Reverse Charge"), (5, 9, 6, 9, "Rate(%)"), (5, 10, 6, 10, "Taxable Value (₹)"),
           (5, 11, 5, 14, "Tax Amount"), (6, 11, 6, 11, "Integrated Tax(₹)"), (6, 12, 6, 12, "Central Tax(₹)"),
           (6, 13, 6, 13, "State/UT Tax(₹)"), (6, 14, 6, 14, "Cess(₹)"), (5, 15, 6, 15, "GSTR-1/5 Period"),
           (5, 16, 6, 16, "GSTR-1/5 Filing Date"), (5, 17, 6, 17, "ITC Availability"), (5, 18, 6, 18, "Reason"),
           (5, 19, 6, 19, "Applicable % of Tax Rate"), (5, 20, 6, 20, "Source"), (5, 21, 6, 21, "IRN"),
           (5, 22, 6, 22, "IRN Date"), (5, 23, 6, 23, "Period")]
    w = {1: 20, 2: 28, 3: 16.9, 4: 15.9, 5: 13, 6: 20, 7: 20, 8: 14, 9: 10, 10: 20, 11: 14, 12: 14, 13: 14, 14: 10,
         15: 20, 16: 16, 17: 14, 18: 35, 19: 19.7, 20: 12, 21: 20, 22: 14, 23: 11.6}
    ws = _sheet(wb, "B2B", "Taxable inward supplies received from registered persons", 23, hdr, 2, w)

    def b2b_row(s, i):
        return [s.get("ctin", ""), s.get("trdnm", "") or s.get("lglNm", ""), i.get("inum", ""),
                INV_TYPES.get(i.get("typ", ""), i.get("typ", "")), _dmy(i.get("dt", "")), _g(i, "val"),
                _pos(i.get("pos")), _yn(i.get("rev")), _rate(i), _g(i, "txval"), _g(i, "igst"), _g(i, "cgst"),
                _g(i, "sgst"), _g(i, "cess"), _prd_long(s.get("supprd")), _dmy(s.get("supfildt", "")),
                ITC_AVL.get(str(i.get("itcavl", "")).upper(), ""), REASONS.get(i.get("rsn", ""), i.get("rsn", "")),
                i.get("diffprcnt", "") if i.get("diffprcnt") not in (None, 1, 1.0) else "",
                "e-Invoice" if i.get("srctyp") else "", i.get("irn", ""), _dmy(i.get("irngendate", "")), period]
    r = 7
    for s, i in sorted(_inv_rows(dd.get("b2b"), period), key=lambda x: (x[0].get("ctin", ""), 0)):
        _put(ws, r, b2b_row(s, i), num_cols=(6, 10, 11, 12, 13, 14), text_cols=(1,))
        r += 1

    # ---- B2BA (revised + original)
    hdr = [(5, 1, 5, 2, "Original Details"), (5, 3, 5, 22, "Revised Details"),
           (6, 1, 7, 1, "Invoice number"), (6, 2, 7, 2, "Invoice Date"), (6, 3, 7, 3, "GSTIN of supplier"),
           (6, 4, 7, 4, "Trade/Legal name"), (6, 5, 6, 8, "Invoice Details"), (7, 5, 7, 5, "Invoice number"),
           (7, 6, 7, 6, "Invoice type"), (7, 7, 7, 7, "Invoice Date"), (7, 8, 7, 8, "Invoice Value(₹)"),
           (6, 9, 7, 9, "Place of supply"), (6, 10, 7, 10, "Supply Attract Reverse Charge"), (6, 11, 7, 11, "Rate(%)"),
           (6, 12, 7, 12, "Taxable Value (₹)"), (6, 13, 6, 16, "Tax Amount"), (7, 13, 7, 13, "Integrated Tax(₹)"),
           (7, 14, 7, 14, "Central Tax(₹)"), (7, 15, 7, 15, "State/UT Tax(₹)"), (7, 16, 7, 16, "Cess(₹)"),
           (6, 17, 7, 17, "GSTR-1/5 Period"), (6, 18, 7, 18, "GSTR-1/5 Filing Date"), (6, 19, 7, 19, "ITC Availability"),
           (6, 20, 7, 20, "Reason"), (6, 21, 7, 21, "Applicable % of Tax Rate"), (6, 22, 7, 22, "Period")]
    ws = _sheet(wb, "B2BA", "Amendments to previously filed invoices by supplier", 22, hdr, 3)
    r = 8
    for s, i in _inv_rows(dd.get("b2ba"), period):
        _put(ws, r, [i.get("oinum", ""), _dmy(i.get("oidt", "")), s.get("ctin", ""), s.get("trdnm", ""),
                     i.get("inum", ""), INV_TYPES.get(i.get("typ", ""), i.get("typ", "")), _dmy(i.get("dt", "")),
                     _g(i, "val"), _pos(i.get("pos")), _yn(i.get("rev")), _rate(i), _g(i, "txval"), _g(i, "igst"),
                     _g(i, "cgst"), _g(i, "sgst"), _g(i, "cess"), _prd_long(s.get("supprd")),
                     _dmy(s.get("supfildt", "")), ITC_AVL.get(str(i.get("itcavl", "")).upper(), ""),
                     REASONS.get(i.get("rsn", ""), i.get("rsn", "")), "", period],
             num_cols=(8, 12, 13, 14, 15, 16))
        r += 1

    # ---- B2B-CDNR
    hdr = [(5, 1, 6, 1, "GSTIN of supplier"), (5, 2, 6, 2, "Trade/Legal name"),
           (5, 3, 5, 7, "Credit note/Debit note details"), (6, 3, 6, 3, "Note number"), (6, 4, 6, 4, "Note type"),
           (6, 5, 6, 5, "Note Supply type"), (6, 6, 6, 6, "Note date"), (6, 7, 6, 7, "Note Value (₹)"),
           (5, 8, 6, 8, "Place of supply"), (5, 9, 6, 9, "Supply Attract Reverse Charge"), (5, 10, 6, 10, "Rate(%)"),
           (5, 11, 6, 11, "Taxable Value (₹)"), (5, 12, 5, 15, "Tax Amount"), (6, 12, 6, 12, "Integrated Tax(₹)"),
           (6, 13, 6, 13, "Central Tax(₹)"), (6, 14, 6, 14, "State/UT Tax(₹)"), (6, 15, 6, 15, "Cess(₹)"),
           (5, 16, 6, 16, "GSTR-1/5 Period"), (5, 17, 6, 17, "GSTR-1/5 Filing Date"), (5, 18, 6, 18, "ITC Availability"),
           (5, 19, 6, 19, "Reason"), (5, 20, 6, 20, "Applicable % of Tax Rate"), (5, 21, 6, 21, "Source"),
           (5, 22, 6, 22, "IRN"), (5, 23, 6, 23, "IRN Date"), (5, 24, 6, 24, "Period")]
    w = {1: 20, 2: 28, 3: 16, 4: 13, 5: 14, 6: 13, 7: 15, 8: 16, 9: 14, 12: 14, 13: 14, 14: 14, 16: 16, 17: 16, 19: 30, 24: 11}
    ws = _sheet(wb, "B2B-CDNR", "Debit/Credit notes (Original)", 24, hdr, 2, w)
    r = 7
    for s, n in _note_rows(dd.get("cdnr")):
        _put(ws, r, [s.get("ctin", ""), s.get("trdnm", "") or s.get("lglNm", ""), n.get("ntnum", ""),
                     NOTE_TYPES.get(n.get("typ", ""), n.get("typ", "")),
                     INV_TYPES.get(n.get("suptyp", ""), "") if n.get("suptyp", "R") != "R" else "",
                     _dmy(n.get("dt", "")), _g(n, "val"), _pos(n.get("pos")), _yn(n.get("rev")), _rate(n),
                     _g(n, "txval"), _g(n, "igst"), _g(n, "cgst"), _g(n, "sgst"), _g(n, "cess"),
                     _prd_long(s.get("supprd")), _dmy(s.get("supfildt", "")),
                     ITC_AVL.get(str(n.get("itcavl", "")).upper(), ""), REASONS.get(n.get("rsn", ""), n.get("rsn", "")),
                     "", "e-Invoice" if n.get("srctyp") else "", n.get("irn", ""), _dmy(n.get("irngendate", "")), period],
             num_cols=(7, 11, 12, 13, 14, 15), text_cols=(1,))
        r += 1

    # ---- B2B-CDNRA
    hdr = [(5, 1, 5, 3, "Original Details"), (5, 4, 5, 24, "Revised Details"),
           (6, 1, 7, 1, "Note type"), (6, 2, 7, 2, "Note number"), (6, 3, 7, 3, "Note date"),
           (6, 4, 7, 4, "GSTIN of supplier"), (6, 5, 7, 5, "Trade/Legal name"),
           (6, 6, 6, 10, "Credit note/Debit note details"), (7, 6, 7, 6, "Note number"), (7, 7, 7, 7, "Note type"),
           (7, 8, 7, 8, "Note Supply type"), (7, 9, 7, 9, "Note date"), (7, 10, 7, 10, "Note Value (₹)"),
           (6, 11, 7, 11, "Place of supply"), (6, 12, 7, 12, "Supply Attract Reverse Charge"), (6, 13, 7, 13, "Rate(%)"),
           (6, 14, 7, 14, "Taxable Value (₹)"), (6, 15, 6, 18, "Tax Amount"), (7, 15, 7, 15, "Integrated Tax(₹)"),
           (7, 16, 7, 16, "Central Tax(₹)"), (7, 17, 7, 17, "State/UT Tax(₹)"), (7, 18, 7, 18, "Cess(₹)"),
           (6, 19, 7, 19, "GSTR-1/5 Period"), (6, 20, 7, 20, "GSTR-1/5 Filing Date"), (6, 21, 7, 21, "ITC Availability"),
           (6, 22, 7, 22, "Reason"), (6, 23, 7, 23, "Applicable % of Tax Rate"), (6, 24, 7, 24, "Period")]
    ws = _sheet(wb, "B2B-CDNRA", "Amendments to previously filed Credit/Debit notes by supplier", 24, hdr, 3)
    r = 8
    for s, n in _note_rows(dd.get("cdnra")):
        _put(ws, r, [NOTE_TYPES.get(n.get("ontyp", n.get("typ", "")), ""), n.get("ontnum", ""), _dmy(n.get("ontdt", "")),
                     s.get("ctin", ""), s.get("trdnm", ""), n.get("ntnum", ""),
                     NOTE_TYPES.get(n.get("typ", ""), n.get("typ", "")), "", _dmy(n.get("dt", "")), _g(n, "val"),
                     _pos(n.get("pos")), _yn(n.get("rev")), _rate(n), _g(n, "txval"), _g(n, "igst"), _g(n, "cgst"),
                     _g(n, "sgst"), _g(n, "cess"), _prd_long(s.get("supprd")), _dmy(s.get("supfildt", "")),
                     ITC_AVL.get(str(n.get("itcavl", "")).upper(), ""), REASONS.get(n.get("rsn", ""), n.get("rsn", "")),
                     "", period],
             num_cols=(10, 14, 15, 16, 17, 18))
        r += 1

    # ---- ISD
    hdr = [(5, 1, 6, 1, "GSTIN of ISD"), (5, 2, 6, 2, "Trade/Legal name"), (5, 3, 6, 3, "ISD Document type"),
           (5, 4, 6, 4, "ISD Document number"), (5, 5, 6, 5, "ISD Document date"), (5, 6, 6, 6, "Original Invoice Number"),
           (5, 7, 6, 7, "Original invoice date"), (5, 8, 5, 11, "Input tax distribution by ISD"),
           (6, 8, 6, 8, "Integrated Tax(₹)"), (6, 9, 6, 9, "Central Tax(₹)"), (6, 10, 6, 10, "State/UT Tax(₹)"),
           (6, 11, 6, 11, "Cess(₹)"), (5, 12, 6, 12, "ISD GSTR-6 Period"), (5, 13, 6, 13, "ISD GSTR-6 Filing Date"),
           (5, 14, 6, 14, "Eligibility of ITC"), (5, 15, 6, 15, "Period")]
    ws = _sheet(wb, "ISD", "ISD Credits", 15, hdr, 2)
    r = 7
    for s in dd.get("isd", []) or []:
        for d in s.get("doclist", []) or []:
            _put(ws, r, [s.get("ctin", ""), s.get("trdnm", ""), d.get("doctyp", ""), d.get("docnum", ""),
                         _dmy(d.get("docdt", "")), d.get("oinvnum", ""), _dmy(d.get("oinvdt", "")), _g(d, "igst"),
                         _g(d, "cgst"), _g(d, "sgst"), _g(d, "cess"), _prd_long(s.get("supprd")),
                         _dmy(s.get("supfildt", "")), ITC_AVL.get(str(d.get("itcelg", "")).upper(), ""), period],
                 num_cols=(8, 9, 10, 11))
            r += 1

    # ---- ISDA
    hdr = [(5, 1, 5, 3, "Original Details"), (5, 4, 5, 18, "Revised Details"),
           (6, 1, 7, 1, "ISD Document type"), (6, 2, 7, 2, "Document Number"), (6, 3, 7, 3, "Document date"),
           (6, 4, 7, 4, "GSTIN of ISD"), (6, 5, 7, 5, "Trade/Legal name"), (6, 6, 7, 6, "ISD Document type"),
           (6, 7, 7, 7, "ISD Document number"), (6, 8, 7, 8, "ISD Document date"), (6, 9, 7, 9, "Original Invoice Number"),
           (6, 10, 7, 10, "Original invoice date"), (6, 11, 6, 14, "Input tax distribution by ISD"),
           (7, 11, 7, 11, "Integrated Tax(₹)"), (7, 12, 7, 12, "Central Tax(₹)"), (7, 13, 7, 13, "State/UT Tax(₹)"),
           (7, 14, 7, 14, "Cess(₹)"), (6, 15, 7, 15, "ISD GSTR-6 Period"), (6, 16, 7, 16, "ISD GSTR-6 Filing Date"),
           (6, 17, 7, 17, "Eligibility of ITC"), (6, 18, 7, 18, "Period")]
    _sheet(wb, "ISDA", "Amendments ISD Credits received", 18, hdr, 3)

    # ---- IMPG
    hdr = [(5, 1, 6, 1, "Icegate Reference Date"), (5, 2, 6, 2, "Port Code"), (5, 3, 5, 5, "Bill of Entry Details"),
           (6, 3, 6, 3, "Number"), (6, 4, 6, 4, "Date"), (6, 5, 6, 5, "Taxable Value"),
           (5, 6, 5, 7, "Amount of tax (₹)"), (6, 6, 6, 6, "Integrated Tax(₹)"), (6, 7, 6, 7, "Cess(₹)"),
           (5, 8, 6, 8, "Amended (Yes)"), (5, 9, 6, 9, "Period")]
    ws = _sheet(wb, "IMPG", "Import of goods from overseas on bill of entry", 9, hdr, 2)
    r = 7
    for b in dd.get("impg", []) or []:
        _put(ws, r, [_dmy(b.get("refdt", "")), b.get("portcode", ""), b.get("boenum", ""), _dmy(b.get("boedt", "")),
                     _g(b, "txval"), _g(b, "igst"), _g(b, "cess"), "Yes" if b.get("isamd") == "Y" else "", period],
             num_cols=(5, 6, 7))
        r += 1

    # ---- IMPGSEZ
    hdr = [(5, 1, 6, 1, "GSTIN of supplier"), (5, 2, 6, 2, "Trade/Legal name"), (5, 3, 6, 3, "Icegate Reference Date"),
           (5, 4, 6, 4, "Port Code"), (5, 5, 5, 7, "Bill of Entry Details"), (6, 5, 6, 5, "Number"),
           (6, 6, 6, 6, "Date"), (6, 7, 6, 7, "Taxable Value"), (5, 8, 5, 9, "Amount of tax (₹)"),
           (6, 8, 6, 8, "Integrated Tax(₹)"), (6, 9, 6, 9, "Cess(₹)"), (5, 10, 6, 10, "Amended (Yes)"),
           (5, 11, 6, 11, "Period")]
    ws = _sheet(wb, "IMPGSEZ", "Import of goods from SEZ units/developers on bill of entry", 11, hdr, 2)
    r = 7
    for s in dd.get("impgsez", []) or []:
        for b in s.get("boe", []) or [s]:
            _put(ws, r, [s.get("ctin", ""), s.get("trdnm", ""), _dmy(b.get("refdt", "")), b.get("portcode", ""),
                         b.get("boenum", ""), _dmy(b.get("boedt", "")), _g(b, "txval"), _g(b, "igst"), _g(b, "cess"),
                         "Yes" if b.get("isamd") == "Y" else "", period], num_cols=(7, 8, 9))
            r += 1

    # ---- Ecomm
    hdr = [(5, 1, 6, 1, "GSTIN of ECO"), (5, 2, 6, 2, "Trade/Legal name"), (5, 3, 5, 6, "Document details"),
           (6, 3, 6, 3, "Document number"), (6, 4, 6, 4, "Document type"), (6, 5, 6, 5, "Document date"),
           (6, 6, 6, 6, "Document value(₹)"), (5, 7, 6, 7, "Place of supply"), (5, 8, 6, 8, "Rate(%)"),
           (5, 9, 6, 9, "Taxable value (₹)"), (5, 10, 5, 13, "Tax amount"), (6, 10, 6, 10, "Integrated Tax(₹)"),
           (6, 11, 6, 11, "Central Tax(₹)"), (6, 12, 6, 12, "State/UT Tax(₹)"), (6, 13, 6, 13, "Cess(₹)"),
           (5, 14, 6, 14, "GSTR-1/IFF/GSTR-1A period"), (5, 15, 6, 15, "GSTR-1/IFF/GSTR-1A filing date"),
           (5, 16, 6, 16, "ITC availability"), (5, 17, 6, 17, "Reason"), (5, 18, 6, 18, "Source"),
           (5, 19, 6, 19, "IRN"), (5, 20, 6, 20, "IRN Date"), (5, 21, 6, 21, "Period")]
    ws = _sheet(wb, "Ecomm", "Documents reported by ECO on which ECO is liable to pay tax u/s 9(5)", 21, hdr, 2)
    r = 7
    for s in dd.get("ecom", []) or []:
        for d in (s.get("inv", []) or []) + (s.get("nt", []) or []):
            if _is_rej(d):
                continue
            _put(ws, r, [s.get("ctin", ""), s.get("trdnm", ""), d.get("inum", d.get("ntnum", "")),
                         INV_TYPES.get(d.get("typ", ""), NOTE_TYPES.get(d.get("typ", ""), d.get("typ", ""))),
                         _dmy(d.get("dt", "")), _g(d, "val"), _pos(d.get("pos")), _rate(d), _g(d, "txval"),
                         _g(d, "igst"), _g(d, "cgst"), _g(d, "sgst"), _g(d, "cess"), _prd_long(s.get("supprd")),
                         _dmy(s.get("supfildt", "")), ITC_AVL.get(str(d.get("itcavl", "")).upper(), ""),
                         REASONS.get(d.get("rsn", ""), d.get("rsn", "")), "e-Invoice" if d.get("srctyp") else "",
                         d.get("irn", ""), _dmy(d.get("irngendate", "")), period], num_cols=(6, 9, 10, 11, 12, 13))
            r += 1

    # ---- Rejected (IMS) sheets
    hdr = [(5, 1, 6, 1, "GSTIN of supplier"), (5, 2, 6, 2, "Trade/Legal name"), (5, 3, 5, 6, "Invoice Details"),
           (6, 3, 6, 3, "Invoice number"), (6, 4, 6, 4, "Invoice type"), (6, 5, 6, 5, "Invoice Date"),
           (6, 6, 6, 6, "Invoice Value(₹)"), (5, 7, 6, 7, "Place of supply"), (5, 8, 6, 8, "Taxable Value (₹)"),
           (5, 9, 5, 12, "Tax Amount"), (6, 9, 6, 9, "Integrated Tax(₹)"), (6, 10, 6, 10, "Central Tax(₹)"),
           (6, 11, 6, 11, "State/UT Tax(₹)"), (6, 12, 6, 12, "Cess(₹)"), (5, 13, 6, 13, "GSTR-1/IFF/GSTR-5 Period"),
           (5, 14, 6, 14, "GSTR-1/IFF/GSTR-5 Filing Date"), (5, 15, 6, 15, "Applicable % of Tax Rate"),
           (5, 16, 6, 16, "Source"), (5, 17, 6, 17, "IRN"), (5, 18, 6, 18, "IRN Date")]
    ws = _sheet(wb, "B2B(Rejected)", "ITC Rejected for taxable inward supplies received from registered persons", 18, hdr, 2)
    r = 7
    for s, i in _inv_rows(dd.get("b2b"), period, rejected=True):
        _put(ws, r, [s.get("ctin", ""), s.get("trdnm", ""), i.get("inum", ""),
                     INV_TYPES.get(i.get("typ", ""), i.get("typ", "")), _dmy(i.get("dt", "")), _g(i, "val"),
                     _pos(i.get("pos")), _g(i, "txval"), _g(i, "igst"), _g(i, "cgst"), _g(i, "sgst"), _g(i, "cess"),
                     _prd_long(s.get("supprd")), _dmy(s.get("supfildt", "")), "", "e-Invoice" if i.get("srctyp") else "",
                     i.get("irn", ""), _dmy(i.get("irngendate", ""))], num_cols=(6, 8, 9, 10, 11, 12), text_cols=(1,))
        r += 1

    hdr = [(5, 1, 6, 1, "GSTIN of supplier"), (5, 2, 6, 2, "Trade/Legal name"),
           (5, 3, 5, 7, "Credit note/Debit note details"), (6, 3, 6, 3, "Note number"), (6, 4, 6, 4, "Note type"),
           (6, 5, 6, 5, "Note Supply type"), (6, 6, 6, 6, "Note date"), (6, 7, 6, 7, "Note Value (₹)"),
           (5, 8, 6, 8, "Place of supply"), (5, 9, 6, 9, "Taxable Value (₹)"), (5, 10, 5, 13, "Tax Amount"),
           (6, 10, 6, 10, "Integrated Tax(₹)"), (6, 11, 6, 11, "Central Tax(₹)"), (6, 12, 6, 12, "State/UT Tax(₹)"),
           (6, 13, 6, 13, "Cess(₹)"), (5, 14, 6, 14, "GSTR-1/IFF/GSTR-5 Period"),
           (5, 15, 6, 15, "GSTR-1/IFF/GSTR-5 Filing Date"), (5, 16, 6, 16, "Applicable % of Tax Rate"),
           (5, 17, 6, 17, "Source"), (5, 18, 6, 18, "IRN"), (5, 19, 6, 19, "IRN Date")]
    ws = _sheet(wb, "B2B-CDNR(Rejected)", "ITC Rejected for Debit/Credit notes (Original)", 19, hdr, 2)
    r = 7
    for s, n in _note_rows(dd.get("cdnr"), rejected=True):
        _put(ws, r, [s.get("ctin", ""), s.get("trdnm", ""), n.get("ntnum", ""),
                     NOTE_TYPES.get(n.get("typ", ""), n.get("typ", "")), "", _dmy(n.get("dt", "")), _g(n, "val"),
                     _pos(n.get("pos")), _g(n, "txval"), _g(n, "igst"), _g(n, "cgst"), _g(n, "sgst"), _g(n, "cess"),
                     _prd_long(s.get("supprd")), _dmy(s.get("supfildt", "")), "", "e-Invoice" if n.get("srctyp") else "",
                     n.get("irn", ""), _dmy(n.get("irngendate", ""))], num_cols=(7, 9, 10, 11, 12, 13), text_cols=(1,))
        r += 1
    return wb


def convert_json_to_portal_excel(data, excel_path, quarterly=False):
    import os
    wb = build_portal_workbook(data, quarterly=quarterly)
    os.makedirs(os.path.dirname(os.path.abspath(excel_path)), exist_ok=True)
    wb.save(excel_path)
    return wb


def convert_and_save_period_excel(raw_file_path, excel_path, period_label="", quarterly=False):
    """
    Takes a downloaded GSTR-2B .json (or .zip containing it) and writes a portal-style Excel
    (ITC Available / ITC not available / ITC Rejected / B2B / B2BA / B2B-CDNR / ... sheets).
    Returns excel_path.
    """
    if raw_file_path.lower().endswith(".zip"):
        with zipfile.ZipFile(raw_file_path, "r") as z:
            json_files = [f for f in z.namelist() if f.lower().endswith(".json")]
            if not json_files:
                raise RuntimeError("No JSON file found inside downloaded zip.")
            with z.open(json_files[0]) as jf:
                data = json.load(jf)
    else:
        with open(raw_file_path, "r", encoding="utf-8") as f:
            data = json.load(f)
    convert_json_to_portal_excel(data, excel_path, quarterly=quarterly)
    return excel_path


# first data row of each detail sheet (header height differs per sheet)
_DETAIL_FIRST_ROW = {"B2B": 7, "B2BA": 8, "B2B-CDNR": 7, "B2B-CDNRA": 8, "ISD": 7, "ISDA": 8, "IMPG": 7,
                     "IMPGSEZ": 7, "Ecomm": 7, "B2B(Rejected)": 7, "B2B-CDNR(Rejected)": 7}
_SUMMARY_SHEETS = ["ITC Available", "ITC not available", "ITC Rejected"]


def merge_all_gstr2b_periods(period_excel_paths, output_excel_path):
    """
    Merges the per-period portal-style workbooks into one workbook:
      - detail sheets (B2B, B2B-CDNR, ...) are stacked one period after another (the 'Period' column tells which)
      - ITC summary sheets are added up across periods
    """
    from copy import copy
    from openpyxl import load_workbook

    paths = list(period_excel_paths)
    base = load_workbook(paths[0])
    for p in paths[1:]:
        wb = load_workbook(p)
        for name, first in _DETAIL_FIRST_ROW.items():
            if name not in base.sheetnames or name not in wb.sheetnames:
                continue
            src, dst = wb[name], base[name]
            nxt = max(dst.max_row + 1, first)
            while nxt > first and all(dst.cell(nxt - 1, c).value in (None, "") for c in range(1, dst.max_column + 1)):
                nxt -= 1
            for r in range(first, src.max_row + 1):
                if all(src.cell(r, c).value in (None, "") for c in range(1, src.max_column + 1)):
                    continue
                for c in range(1, src.max_column + 1):
                    s, d = src.cell(r, c), dst.cell(nxt, c)
                    d.value = s.value
                    d.font, d.alignment, d.number_format = copy(s.font), copy(s.alignment), s.number_format
                nxt += 1
        for name in _SUMMARY_SHEETS:
            if name not in base.sheetnames or name not in wb.sheetnames:
                continue
            src, dst = wb[name], base[name]
            for r in range(7, src.max_row + 1):
                for c in range(4, 8):
                    sv, dv = src.cell(r, c).value, dst.cell(r, c).value
                    if isinstance(sv, (int, float)) and isinstance(dv, (int, float)) and not type(dst.cell(r, c)).__name__ == "MergedCell":
                        dst.cell(r, c).value = round(dv + sv, 2)
    base["Read me"]["C5"].value = "Merged (multiple periods)"
    os.makedirs(os.path.dirname(os.path.abspath(output_excel_path)), exist_ok=True)
    base.save(output_excel_path)
    return output_excel_path


# ----------------------------------------------------------------------------
# Detection & GSTR-2B Download Routine
# ----------------------------------------------------------------------------
def detect_gst_filing_frequency(driver):
    """
    Checks if taxpayer is registered as Quarterly (QRMP) or Monthly:
      1. Inspects welcome screen text: 'Return filing preference : Quarterly'
      2. Inspects returns dashboard blue advisory banner: 'quarterly frequency'
    """
    try:
        body_text = driver.find_element(By.TAG_NAME, "body").text.lower()
        if "quarterly frequency" in body_text or "return filing preference" in body_text and "quarterly" in body_text:
            return "Quarterly"
        if "monthly frequency" in body_text or "return filing preference" in body_text and "monthly" in body_text:
            return "Monthly"
    except Exception:
        pass
    return "Unknown"


def go_back_to_returns_dashboard(driver):
    """
    From the 'Offline Download for Quarterly GSTR-2B' page, click the BACK button so the
    Returns dashboard (Financial Year / Quarter / Period / SEARCH) is shown again.
    Falls back to opening the dashboard URL directly if BACK cannot be clicked.
    """
    dismiss_alert(driver)
    back_xps = [
        "//button[normalize-space(.)='BACK' or normalize-space(.)='Back']",
        "//a[normalize-space(.)='BACK' or normalize-space(.)='Back']",
        "//input[(@type='button' or @type='submit') and (@value='BACK' or @value='Back')]",
        "//*[self::button or self::a][contains(translate(normalize-space(.), 'back', 'BACK'), 'BACK')]",
    ]
    clicked = click_xpath(driver, back_xps, timeout=8)
    if clicked:
        time.sleep(2)
        dismiss_alert(driver)
        dismiss_gst_popups(driver)
        try:
            WebDriverWait(driver, 15).until(
                lambda d: any(e.is_displayed() for e in d.find_elements(By.TAG_NAME, "select")))
            return True
        except TimeoutException:
            pass

    # Fallback: open the returns dashboard directly
    try:
        driver.get(GST_RETURNS_URL)
        WebDriverWait(driver, 30).until(
            lambda d: any(e.is_displayed() for e in d.find_elements(By.TAG_NAME, "select")))
        dismiss_gst_popups(driver)
        return True
    except Exception:
        return False


def on_offline_download_page(driver):
    """True if the browser is currently on the 'Offline Download for Quarterly GSTR-2B' page."""
    try:
        return bool(driver.find_elements(
            By.XPATH, "//*[contains(normalize-space(.), 'Offline Download for')]"
        )) and not any(e.is_displayed() for e in driver.find_elements(By.TAG_NAME, "select"))
    except Exception:
        return False



def navigate_to_returns_dashboard(driver):
    dismiss_gst_popups(driver)
    if on_offline_download_page(driver):
        go_back_to_returns_dashboard(driver)
        return
    ret_btn_xps = [
        "//button[contains(.,'RETURN DASHBOARD') or contains(.,'Return Dashboard')]",
        "//a[contains(.,'RETURN DASHBOARD') or contains(.,'Return Dashboard')]",
        "//button[contains(.,'FILE RETURNS') or contains(.,'File Returns')]"
    ]
    if click_xpath(driver, ret_btn_xps, timeout=4):
        time.sleep(2)
        dismiss_gst_popups(driver)
        return

    if "return.gst.gov.in" not in driver.current_url.lower():
        driver.get(GST_RETURNS_URL)
        time.sleep(2)
    dismiss_gst_popups(driver)


def gst_download_gstr2b_json(driver, month, year, temp_dir, is_quarterly=False):
    """
    Searches for the period on the Returns Dashboard, clicks DOWNLOAD on the GSTR-2B tile,
    and clicks 'GENERATE JSON FILE TO DOWNLOAD' on the offline download page.
    Returns: (status, filename_or_error_info)
    """
    m_no = MONTHS.index(month) + 1
    fy_start = int(year) if m_no >= 4 else int(year) - 1
    fy_text = f"{fy_start}-{str(fy_start + 1)[2:]}"
    quarter_num = ((m_no - 4) % 12) // 3 + 1

    navigate_to_returns_dashboard(driver)
    WebDriverWait(driver, 40).until(
        lambda d: d.find_elements(By.TAG_NAME, "select") or "login" in d.current_url.lower()
    )
    if "login" in driver.current_url.lower():
        return "session_expired", "Session expired"
    time.sleep(1.5)

    # Select FY, Quarter, Period
    gst_select(driver, ["fin"], 0, fy_text)
    gst_select(driver, ["quarter"], 1, f"Quarter {quarter_num}", starts=True)
    gst_select(driver, ["mon"], 2, month)

    # Click SEARCH
    if not click_xpath(driver, ["//button[contains(normalize-space(.), 'SEARCH') or contains(.,'Search')]"], timeout=10):
        return "error", "SEARCH button not found"
    time.sleep(2.5)

    # Check if GSTR-2B tile is present
    tile_xps = ["//*[contains(text(),'GSTR-2B') or contains(text(),'GSTR2B')]"]
    tiles = driver.find_elements(By.XPATH, tile_xps[0])
    if not tiles:
        if is_quarterly and m_no not in [3, 6, 9, 12]:
            return "only_2a", f"GSTR-2B not present for {month} {year} (Quarterly filer: available in quarter-end month)"
        return "not_available", f"GSTR-2B tile not found for {month} {year}"

    # Click 'DOWNLOAD' button on the GSTR-2B tile
    dl_tile_xps = [
        "//*[contains(text(),'GSTR-2B') or contains(text(),'GSTR2B')]/ancestor::div[.//button[contains(translate(.,'DOWNLOAD','download'),'download')]][1]//button[contains(translate(.,'DOWNLOAD','download'),'download')]",
        "//*[contains(text(),'GSTR-2B') or contains(text(),'GSTR2B')]/ancestor::div[.//button[contains(translate(.,'DOWNLOAD','download'),'download')]][2]//button[contains(translate(.,'DOWNLOAD','download'),'download')]",
        "//button[contains(.,'DOWNLOAD') and (contains(@onclick,'2B') or contains(@id,'2B'))]"
    ]

    clicked_tile_dl = click_xpath(driver, dl_tile_xps, timeout=10)
    if not clicked_tile_dl:
        view_xps = [
            "//*[contains(text(),'GSTR-2B') or contains(text(),'GSTR2B')]/ancestor::div[.//button[contains(translate(.,'VIEW','view'),'view')]][1]//button[contains(translate(.,'VIEW','view'),'view')]"
        ]
        click_xpath(driver, view_xps, timeout=10)

    time.sleep(3)
    dismiss_gst_popups(driver)

    # Offline Download page
    before = list_files(temp_dir)

    json_dl_xps = [
        "//button[contains(normalize-space(.), 'GENERATE JSON FILE TO DOWNLOAD')]",
        "//button[contains(translate(., 'JSON', 'json'), 'json') and contains(translate(., 'GENERATE', 'generate'), 'generate')]",
        "//button[contains(translate(., 'JSON', 'json'), 'json')]",
        "//a[contains(translate(., 'JSON', 'json'), 'json') and contains(translate(., 'DOWNLOAD', 'download'), 'download')]",
        "//*[(self::a or self::button) and contains(., 'Click here to download')]"
    ]

    if not click_xpath(driver, json_dl_xps, timeout=15):
        go_back_to_returns_dashboard(driver)
        return "error", "GENERATE JSON FILE TO DOWNLOAD button not found on download page"

    time.sleep(2)
    dismiss_alert(driver)

    result = None
    f = wait_new_file(temp_dir, before, DOWNLOAD_WAIT_SECONDS)
    if f:
        result = ("downloaded", f)
    else:
        click_here_xps = [
            "//a[contains(translate(., 'DOWNLOAD', 'download'), 'download') or contains(., 'Click here to download')]",
            "//*[contains(text(), 'Click here to download')]"
        ]
        if click_xpath(driver, click_here_xps, timeout=10):
            f = wait_new_file(temp_dir, before, DOWNLOAD_WAIT_SECONDS)
            if f:
                result = ("downloaded", f)
    if result is None:
        result = ("timeout", "JSON file was not received within timeout.")

    # Press BACK so the Returns dashboard (FY / Quarter / Period / SEARCH) is shown for the next period
    go_back_to_returns_dashboard(driver)
    return result


def run_gstr2b(cfg, temp_dir):
    driver = make_driver(temp_dir)
    user = cfg["username"]
    log_rows = []
    parsed_period_dfs = []
    saved_json_files = []

    try:
        gst_login(driver, user, cfg["password"], fill_pw=not cfg["manual_pw"])
        if not wait_for_login(driver, home_part="/auth/"):
            print("Login not completed in time. Exiting.")
            return
        print("Logged in successfully to GST Portal.")

        dismiss_gst_popups(driver)
        time.sleep(1)

        is_quarterly = (cfg["frequency"] == FREQ_Q)
        print(f"Client type selected: {'Quarterly (QRMP)' if is_quarterly else 'Monthly'}")
        try:
            portal_says = detect_gst_filing_frequency(driver)
            if portal_says in ("Quarterly", "Monthly") and (portal_says == "Quarterly") != is_quarterly:
                print(f"   NOTE: portal text suggests this GSTIN is {portal_says}, but you selected "
                      f"{'Quarterly' if is_quarterly else 'Monthly'}. Continuing with your selection.")
        except Exception:
            pass

        months_to_download = cfg["months"]
        if is_quarterly:
            to_dl, qinfo, warnings, skipped = resolve_quarterly_periods(cfg["months"])
            print(f"Smart Quarterly Plan:")
            for m, y, q_end in skipped:
                print(f"  • {m} {y}: Skipped (GSTR-2B generated quarterly in {q_end[0]} {q_end[1]})")

            if warnings:
                w_lines = []
                missing_months_to_add = []
                for w in warnings:
                    sel_str = " and ".join([f"{m} {y}" for m, y in w["selected"]])
                    q_end_str = f"{w['missing_q_end'][0]} {w['missing_q_end'][1]}"
                    w_lines.append(f"• {sel_str} cannot be downloaded because GSTR-2B is generated quarterly.\n"
                                   f"  For downloading data for this quarter, please select {q_end_str}.")
                    missing_months_to_add.append(w["missing_q_end"])
                    for sm, sy in w["selected"]:
                        log_rows.append([user, "GSTR-2B", sm, sy, "SKIPPED",
                                        f"Cannot download without {q_end_str} (Quarterly GSTR-2B)"])

                warning_text = "\n\n".join(w_lines)
                print("\n[WARNING]\n" + warning_text)

                try:
                    r = tk.Tk()
                    r.withdraw()
                    r.attributes("-topmost", True)
                    opt = messagebox.askyesno(
                        "Quarterly GSTR-2B Warning",
                        f"{warning_text}\n\n"
                        f"Would you like to include {', '.join([f'{m} {y}' for m, y in missing_months_to_add])} "
                        f"now so that the quarter's GSTR-2B can be downloaded?"
                    )
                    r.destroy()
                except Exception:
                    opt = False

                if opt:
                    for mm in missing_months_to_add:
                        if mm not in to_dl:
                            to_dl.append(mm)
                    to_dl = sorted(to_dl, key=lambda x: (int(x[1]), MONTHS.index(x[0])))

            months_to_download = to_dl

        total = len(months_to_download)
        if total == 0:
            print("No valid quarter-end periods selected for download.")
            return

        for n, (month, year) in enumerate(months_to_download, 1):
            label = f"GSTR-2B {month} {year}"
            print(f"\n[{n}/{total}] Processing {label}...")
            try:
                status, info = gst_download_gstr2b_json(driver, month, year, temp_dir, is_quarterly=is_quarterly)
            except Exception as e:
                status, info = "error", str(e)[:200]

            if status == "downloaded":
                raw_file = info
                m_no = MONTHS.index(month) + 1
                period_tag = f"{year}-{m_no:02d}"

                # Save raw JSON file
                json_dest_dir = os.path.join(cfg["out"], user, "GSTR2B", "JSON")
                os.makedirs(json_dest_dir, exist_ok=True)
                json_dest = os.path.join(json_dest_dir, raw_file)
                shutil.copy2(os.path.join(temp_dir, raw_file), json_dest)
                saved_json_files.append(json_dest)

                # Convert JSON to Excel (.xlsx)
                excel_dest_dir = os.path.join(cfg["out"], user, "GSTR2B")
                os.makedirs(excel_dest_dir, exist_ok=True)
                excel_dest = os.path.join(excel_dest_dir, f"{user}_GSTR2B_{period_tag}.xlsx")

                try:
                    convert_and_save_period_excel(
                        os.path.join(temp_dir, raw_file),
                        excel_dest,
                        period_label=f"{month[:3]}-{year}",
                        quarterly=is_quarterly
                    )
                    parsed_period_dfs.append(excel_dest)
                    print(f"[{n}/{total}] {label}: Saved JSON and converted to Excel -> {excel_dest}")
                except Exception as e:
                    print(f"[{n}/{total}] {label}: JSON downloaded, but Excel conversion failed: {e}")

                log_rows.append([user, "GSTR-2B", month, year, "downloaded", excel_dest])
            elif status == "session_expired":
                print(f"[{n}/{total}] {label}: SESSION EXPIRED. Please login again.")
                log_rows.append([user, "GSTR-2B", month, year, status, "Session expired"])
                break
            elif status == "only_2a":
                print(f"[{n}/{total}] {label}: {info}")
                log_rows.append([user, "GSTR-2B", month, year, "only_gstr2a", info])
            else:
                print(f"[{n}/{total}] {label}: {status.upper()} - {info}")
                log_rows.append([user, "GSTR-2B", month, year, status, info or ""])

            time.sleep(1)

        # Merge all downloaded periods into a single master Excel file
        if parsed_period_dfs:
            print("\nMerging all downloaded GSTR-2B periods into master Excel...")
            first_p = months_to_download[0]
            last_p = months_to_download[-1]
            rng_tag = f"{first_p[0][:3]}{first_p[1]}-{last_p[0][:3]}{last_p[1]}"
            merged_excel_dest = os.path.join(
                cfg["out"], user, f"{user}_GSTR2B_MERGED_{rng_tag}.xlsx"
            )
            try:
                merge_all_gstr2b_periods(parsed_period_dfs, merged_excel_dest)
                print(f"Master Merged File created: {merged_excel_dest}")
            except Exception as e:
                print(f"Merge error: {e}")

    finally:
        try:
            driver.quit()
        except Exception:
            pass
        shutil.rmtree(temp_dir, ignore_errors=True)

        if log_rows:
            log_path = os.path.join(cfg["out"], "download_log.csv")
            new = not os.path.exists(log_path)
            with open(log_path, "a", newline="", encoding="utf-8") as f:
                w = csv.writer(f)
                if new:
                    w.writerow(["GSTIN", "Type", "Month", "Year", "Status", "Info"])
                w.writerows(log_rows)

    ok = sum(1 for r in log_rows if r[4] == "downloaded")
    sk = sum(1 for r in log_rows if r[4] in ["SKIPPED", "only_gstr2a"])
    er = len(log_rows) - ok - sk
    print(f"\nDONE. Downloaded: {ok} | Skipped/QRMP: {sk} | Errors: {er}")
    print(f"Folder: {os.path.join(cfg['out'], user, 'GSTR2B')}")

    try:
        r = tk.Tk()
        r.withdraw()
        messagebox.showinfo(
            "Download Summary",
            f"GSTR-2B Download Complete!\n\n"
            f"Downloaded: {ok} periods\n"
            f"Skipped / Quarterly handled: {sk}\n"
            f"Errors / Timeouts: {er}\n\n"
            f"Saved to: {os.path.join(cfg['out'], user, 'GSTR2B')}"
        )
        r.destroy()
    except Exception:
        pass


# ----------------------------------------------------------------------------
# E-Way Bill Download Flow (Preserved)
# ----------------------------------------------------------------------------
def wait_selects(driver):
    WebDriverWait(driver, 20).until(
        lambda d: len(d.find_elements(By.TAG_NAME, "select")) >= 2)


def set_select(driver, index, text):
    wait_selects(driver)
    el = driver.find_elements(By.TAG_NAME, "select")[index]
    sel = Select(el)
    if sel.first_selected_option.text.strip() == text:
        return
    sel.select_by_visible_text(text)
    try:
        WebDriverWait(driver, 4).until(EC.staleness_of(el))
    except TimeoutException:
        pass
    wait_selects(driver)
    time.sleep(0.5)


def set_radio(driver, dtype):
    wait_selects(driver)
    radios = driver.find_elements(By.CSS_SELECTOR, "input[type='radio']")
    if len(radios) < 2:
        return
    r = radios[0 if dtype == "Outward" else 1]
    if r.is_selected():
        return
    driver.execute_script("arguments[0].click();", r)
    try:
        WebDriverWait(driver, 4).until(EC.staleness_of(r))
    except TimeoutException:
        pass
    wait_selects(driver)
    time.sleep(0.5)


def download_ewb_month(driver, dtype, month, year, temp_dir):
    driver.get(REPORT_URL)
    if "login" in driver.current_url.lower():
        return "session_expired", None

    wait_selects(driver)
    set_radio(driver, dtype)
    set_select(driver, 0, month)
    set_select(driver, 1, year)

    go = first_present(driver, [
        (By.XPATH, "//input[(@type='submit' or @type='button') and @value='Go']"),
        (By.XPATH, "//button[normalize-space()='Go']"),
    ])
    if go is None:
        return "error", "Go button not found"

    before = list_files(temp_dir)
    driver.execute_script("arguments[0].click();", go)
    time.sleep(1.5)
    alert_msg = dismiss_alert(driver)

    f = wait_new_file(temp_dir, before, DOWNLOAD_WAIT_SECONDS)
    if f:
        return "downloaded", f
    return "no_data", alert_msg


def promote_header(df):
    df = df.fillna("").astype(str)
    filled = df.apply(lambda c: c.str.strip() != "")
    cnt = filled.sum(axis=1)
    if df.shape[1] > 1:
        same = df.apply(lambda r: len({v.strip() for v in r if v.strip()}) <= 1, axis=1)
        cnt = cnt.where(~same, 0)
    head = cnt.head(15)
    if head.empty or head.max() == 0:
        return pd.DataFrame()
    thr = head.max() * 0.6
    h = next(i for i, v in enumerate(head) if v >= thr)
    cols, seen = [], {}
    for i, v in enumerate(df.iloc[h]):
        v = v.strip() or f"Col{i + 1}"
        if v in seen:
            seen[v] += 1
            v = f"{v}_{seen[v]}"
        else:
            seen[v] = 0
        cols.append(v)
    out = df.iloc[h + 1:].copy()
    out.columns = cols
    out = out[filled.iloc[h + 1:].any(axis=1)]
    return out.reset_index(drop=True)


def read_any(path):
    errs = []
    try:
        raw = pd.read_excel(path, sheet_name=None, header=None, dtype=str)
        return {n: promote_header(d) for n, d in raw.items()}
    except Exception as e:
        errs.append("excel: " + str(e)[:100])
    try:
        tables = pd.read_html(path, header=None, keep_default_na=False)
        return {f"Sheet{i + 1}": promote_header(t) for i, t in enumerate(tables)}
    except Exception as e:
        errs.append("html: " + str(e)[:100])
    raise RuntimeError(" | ".join(errs))


def smart_numeric(df):
    for col in df.columns:
        s = df[col].astype(str).str.strip()
        nonblank = s[s != ""]
        if nonblank.empty:
            continue
        if nonblank.str.fullmatch(r"-?\d+(\.\d+)?").all():
            if nonblank.str.fullmatch(r"\d{10,}").any() or nonblank.str.fullmatch(r"0\d+").any():
                continue
            df[col] = pd.to_numeric(s.where(s != ""), errors="coerce")
    return df


def merge_ewb_files(saved_files, gstin, cfg):
    first, last = cfg["months"][0], cfg["months"][-1]
    rng = f"{first[0][:3]}{first[1]}-{last[0][:3]}{last[1]}"
    outputs = []

    for dtype in cfg["types"]:
        items = sorted([x for x in saved_files if x[0] == dtype],
                       key=lambda x: (int(x[1]), x[2]))
        if not items:
            continue
        merged = {}
        for _, year, m_no, month, path in items:
            try:
                sheets = read_any(path)
            except Exception as e:
                print(f"   Merge: could not read {os.path.basename(path)} ({e})")
                continue
            for name, df in sheets.items():
                if df.empty:
                    continue
                df.insert(0, "Source_Type", dtype)
                df.insert(0, "Source_Period", f"{month[:3]}-{year}")
                merged.setdefault(name, []).append(df)
        if not merged:
            continue

        dest_dir = os.path.join(cfg["out"], gstin)
        os.makedirs(dest_dir, exist_ok=True)
        dest = os.path.join(dest_dir, f"{gstin}_{dtype}_MERGED_{rng}.xlsx")
        if os.path.exists(dest):
            try:
                os.remove(dest)
            except PermissionError:
                dest = dest.replace(".xlsx", f"_{datetime.now():%H%M%S}.xlsx")
        try:
            with pd.ExcelWriter(dest, engine="openpyxl") as w:
                for name, dfs in merged.items():
                    big = pd.concat(dfs, ignore_index=True, sort=False).fillna("")
                    big = smart_numeric(big)
                    big.to_excel(w, sheet_name=str(name)[:31] or "Sheet1", index=False)
            print(f"Merged {dtype}: {len(items)} files -> {dest}")
            outputs.append(dest)
        except Exception as e:
            print(f"   Merge: could not write {dtype} merged file ({e})")
    return outputs


def run_ewb(cfg, temp_dir):
    driver = make_driver(temp_dir)
    log_rows = []
    saved_files = []
    gstin = cfg["username"]
    try:
        fill_login(driver, cfg["username"], cfg["password"], fill_pw=not cfg["manual_pw"])
        refill_fn = None if cfg["manual_pw"] else (
            lambda: fill_credentials(driver, cfg["username"], cfg["password"], True))
        if not wait_for_login(driver, refill=refill_fn):
            print("Login not completed in time. Exiting.")
            return

        gstin = get_gstin(driver, cfg["username"])
        print(f"Logged in. GSTIN: {gstin}")

        total = len(cfg["types"]) * len(cfg["months"])
        n = 0
        stop = False
        for dtype in cfg["types"]:
            if stop:
                break
            for month, year in cfg["months"]:
                n += 1
                label = f"{dtype} {month} {year}"
                try:
                    status, info = download_ewb_month(driver, dtype, month, year, temp_dir)
                except Exception as e:
                    status, info = "error", str(e)[:120]

                if status == "downloaded":
                    ext = os.path.splitext(info)[1] or ".xls"
                    m_no = MONTHS.index(month) + 1
                    dest_dir = os.path.join(cfg["out"], gstin, dtype)
                    os.makedirs(dest_dir, exist_ok=True)
                    dest = os.path.join(dest_dir, f"{gstin}_{dtype}_{year}-{m_no:02d}{ext}")
                    if os.path.exists(dest):
                        try:
                            os.remove(dest)
                        except PermissionError:
                            base, e = os.path.splitext(dest)
                            dest = f"{base}_{datetime.now():%H%M%S}{e}"
                    shutil.move(os.path.join(temp_dir, info), dest)
                    saved_files.append((dtype, year, m_no, month, dest))
                    print(f"[{n}/{total}] {label}: saved")
                elif status == "no_data":
                    print(f"[{n}/{total}] {label}: no data")
                elif status == "session_expired":
                    print(f"[{n}/{total}] {label}: SESSION EXPIRED - login again and re-run for remaining months")
                    log_rows.append([gstin, dtype, month, year, status, ""])
                    stop = True
                    break
                else:
                    print(f"[{n}/{total}] {label}: ERROR {info}")
                log_rows.append([gstin, dtype, month, year, status, info or ""])
                time.sleep(1)
    finally:
        try:
            driver.quit()
        except Exception:
            pass
        shutil.rmtree(temp_dir, ignore_errors=True)

        if log_rows:
            log_path = os.path.join(cfg["out"], "download_log.csv")
            new = not os.path.exists(log_path)
            with open(log_path, "a", newline="", encoding="utf-8") as f:
                w = csv.writer(f)
                if new:
                    w.writerow(["GSTIN", "Type", "Month", "Year", "Status", "Info"])
                w.writerows(log_rows)

    merged_paths = []
    if saved_files:
        print("\nMerging files...")
        try:
            merged_paths = merge_ewb_files(saved_files, gstin, cfg)
        except Exception as e:
            print(f"Merge failed: {e}")

    ok = sum(1 for r in log_rows if r[4] == "downloaded")
    nd = sum(1 for r in log_rows if r[4] == "no_data")
    er = len(log_rows) - ok - nd
    print(f"\nDONE. Downloaded: {ok} | No data: {nd} | Errors: {er}")


# ----------------------------------------------------------------------------
# Main entry
# ----------------------------------------------------------------------------
def main():
    cfg = get_inputs()
    if not cfg:
        return

    if cfg["remember"] and cfg["password"]:
        save_user(cfg["username"], cfg["password"], cfg["service"])

    os.makedirs(cfg["out"], exist_ok=True)
    temp_dir = os.path.join(cfg["out"], "_temp_download")
    os.makedirs(temp_dir, exist_ok=True)

    if cfg["service"] == SERVICE_2B:
        run_gstr2b(cfg, temp_dir)
    else:
        run_ewb(cfg, temp_dir)


if __name__ == "__main__":
    try:
        main()
    except Exception:
        import traceback
        traceback.print_exc()
        input("\nAn error occurred. Press Enter to close...")
